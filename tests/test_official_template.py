import os
import openpyxl
from datetime import datetime, date
from decimal import Decimal
from sqlalchemy.orm import Session

from app.models.perfil_regras import PerfilRegras
from app.models.empresa import Empresa
from app.models.regra_aliquota import RegraAliquotaDestino
from app.models.solicitacao import Solicitacao
from app.models.template_xlsx import TemplateXlsx
from app.core.seeds import (
    seed_default_templates,
    DEFAULT_ANTECIPACAO_PARCIAL_TEMPLATE_PATH,
    DEFAULT_ANTECIPACAO_PARCIAL_ANTECIPADO_TEMPLATE_PATH,
    DEFAULT_ANTECIPACAO_TRIBUTARIA_ANTECIPADO_TEMPLATE_PATH,
)
from app.services.pipeline_service import ProcessingPipelineService
from app.services.excel.formula_guard import FormulaGuard


def test_modelo_oficial_predefinido_tributaria_pago_antecipadamente(db_session):
    assert os.path.exists(DEFAULT_ANTECIPACAO_TRIBUTARIA_ANTECIPADO_TEMPLATE_PATH)
    seed_default_templates(db_session)

    template = (
        db_session.query(TemplateXlsx)
        .filter(
            TemplateXlsx.tipo == "antecipacao_tributaria_antecipado",
            TemplateXlsx.ativo == True,
        )
        .one()
    )
    assert template.capacidade_linhas == 35
    assert template.mapeamento_campos["sheet_name"] == "4"
    assert template.mapeamento_campos["columns"]["data_entrada"] == "B"
    assert template.mapeamento_campos["columns"]["mva"] == "H"

    wb = openpyxl.load_workbook(template.arquivo_path, data_only=False)
    assert "4" in wb.sheetnames
    ws = wb["4"]
    assert "PAGOS ANTECIPADAMENTE" in ws["F1"].value
    assert ws["N4"].value == '=IF(L4="N",E4*H4%+E4,(E4+E4*H4%)-((E4+E4*H4%)*I4%))'
    assert ws["O4"].value == "=N4*J4%"
    assert ws["P4"].value == "=R4*K4%"
    assert ws["Q4"].value == "=O4-P4"
    assert ws["T4"].value == "=Q4-S4"
    wb.close()

def test_modelo_oficial_predefinido_antecipacao_parcial(db_session, sample_xml_nfe):
    # 1. Semear template oficial
    assert os.path.exists(DEFAULT_ANTECIPACAO_PARCIAL_TEMPLATE_PATH)
    seed_default_templates(db_session)

    template_ativo = (
        db_session.query(TemplateXlsx)
        .filter(TemplateXlsx.tipo == "antecipacao_parcial", TemplateXlsx.ativo == True)
        .first()
    )
    assert template_ativo is not None
    assert template_ativo.versao >= 1

    # 2. Cadastrar Perfil e Regras
    perfil = PerfilRegras(nome="Perfil Padrão Bahia")
    db_session.add(perfil)
    db_session.commit()

    # Regra padrão estadual BA (20.5%)
    db_session.add(RegraAliquotaDestino(
        perfil_regras_id=perfil.id,
        uf="BA",
        ncm=None,
        aliquota=Decimal("0.2050"),
        descricao="Alíquota com FECOP BA"
    ))
    db_session.commit()

    # 3. Cadastrar Empresa
    empresa = Empresa(
        razao_social="CAMATEX CONFECCOES MATRIZ",
        cnpj="12345678000195",
        uf="BA",
        perfil_regras_id=perfil.id
    )
    db_session.add(empresa)
    db_session.commit()

    # 4. Criar Solicitação de Processamento para Janeiro/2026
    solicitacao = Solicitacao(
        empresa_id=empresa.id,
        periodo_inicio=date(2026, 1, 1),
        periodo_fim=date(2026, 1, 31),
        tipo_planilha="antecipacao_parcial",
        template_id=template_ativo.id,
        status="pendente"
    )
    db_session.add(solicitacao)
    db_session.commit()

    # 5. Executar Pipeline
    service = ProcessingPipelineService(db_session)
    solicitacao_processada = service.process_solicitacao(
        solicitacao_id=solicitacao.id,
        xml_files_bytes=[("nfe_1234.xml", sample_xml_nfe)]
    )

    assert solicitacao_processada.status == "concluido"
    assert solicitacao_processada.total_notas_processadas == 1
    assert os.path.exists(solicitacao_processada.arquivo_saida_path)

    # 6. Validar o arquivo gerado
    wb = openpyxl.load_workbook(solicitacao_processada.arquivo_saida_path, data_only=False)
    
    # Aba JANEIRO deve ter sido selecionada
    assert "JANEIRO" in wb.sheetnames
    ws = wb["JANEIRO"]

    # Header da empresa preenchido
    assert "CAMATEX CONFECCOES MATRIZ" in ws["A2"].value
    assert "01/2026" in ws["A2"].value

    # Linha 4 preenchida com os dados da NF-e
    assert ws["A4"].value == 1            # Sequencial
    assert str(ws["D4"].value) == "1234"  # N. Fiscal
    assert float(ws["E4"].value) == 5400.0 # V. Total
    assert float(ws["G4"].value) == 400.0  # IPI + Despesas
    assert float(ws["H4"].value) == 20.5   # A. DST (20.5%)
    assert float(ws["I4"].value) == 12.0   # A. ORI (12% do XML)

    # Fórmulas de cálculo intactas na linha 4
    assert ws["F4"].value == "=E4-G4"
    assert ws["J4"].value == "=(E4)/100*H4"
    assert ws["K4"].value == "=(F4)/100*I4"
    assert ws["L4"].value == "=J4-K4"

    # Fórmulas de todas as outras linhas (4 a 75) continuam presentes
    for r in range(4, 76):
        assert ws[f"F{r}"].value == f"=E{r}-G{r}"
        assert ws[f"J{r}"].value == f"=(E{r})/100*H{r}"
        assert ws[f"K{r}"].value == f"=(F{r})/100*I{r}"
        assert ws[f"L{r}"].value == f"=J{r}-K{r}"

    # Totais no rodapé (Linha 76) intactos
    assert ws["E76"].value == "=SUM(E4:E75)"
    assert ws["F76"].value == "=SUM(F4:F75)"
    assert ws["G76"].value == "=SUM(G4:G75)"
    assert ws["J76"].value == "=SUM(J4:J75)"
    assert ws["K76"].value == "=SUM(K4:K75)"
    assert ws["L76"].value == "=SUM(L4:L75)"

    wb.close()


def test_modelo_oficial_predefinido_antecipacao_parcial_antecipado(
        db_session, xml_nf901, xml_nf902, sped_janeiro_com_nf901):
    # 1. Semear template oficial
    assert os.path.exists(DEFAULT_ANTECIPACAO_PARCIAL_ANTECIPADO_TEMPLATE_PATH)
    seed_default_templates(db_session)

    template_ativo = (
        db_session.query(TemplateXlsx)
        .filter(TemplateXlsx.tipo == "antecipacao_parcial_antecipado", TemplateXlsx.ativo == True)
        .first()
    )
    assert template_ativo is not None
    assert template_ativo.versao >= 1

    # 2. Cadastrar Perfil e Regras
    perfil = PerfilRegras(nome="Perfil Padrão Bahia Antecipado")
    db_session.add(perfil)
    db_session.commit()

    db_session.add(RegraAliquotaDestino(
        perfil_regras_id=perfil.id,
        uf="BA",
        ncm=None,
        aliquota=Decimal("0.2050"),
        descricao="Alíquota com FECOP BA"
    ))
    db_session.commit()

    # 3. Cadastrar Empresa
    empresa = Empresa(
        razao_social="CAMATEX CONFECCOES MATRIZ",
        cnpj="12345678000195",
        uf="BA",
        perfil_regras_id=perfil.id
    )
    db_session.add(empresa)
    db_session.commit()

    # 4. Criar Solicitação de Processamento Multi-Saída
    solicitacao = Solicitacao(
        empresa_id=empresa.id,
        periodo_inicio=date(2026, 1, 1),
        periodo_fim=date(2026, 1, 31),
        tipo_planilha=None,
        status="pendente"
    )
    db_session.add(solicitacao)
    db_session.commit()

    # 5. Executar Pipeline
    res = ProcessingPipelineService(db_session).process_solicitacao(
        solicitacao_id=solicitacao.id,
        xml_files_bytes=[("901.xml", xml_nf901), ("902.xml", xml_nf902)],
        sped_file_bytes=sped_janeiro_com_nf901,
    )
    db_session.refresh(res)

    assert res.status == "concluido"
    saidas = {s.tipo: s for s in res.saidas}
    assert "antecipacao_parcial_antecipado" in saidas
    arquivo_path = saidas["antecipacao_parcial_antecipado"].arquivo_path
    assert os.path.exists(arquivo_path)

    # 6. Validar o arquivo gerado (RP-155)
    wb = openpyxl.load_workbook(arquivo_path, data_only=False)
    assert "01-2026" in wb.sheetnames
    ws = wb["01-2026"]

    # Header da empresa preenchido em A2
    assert "CAMATEX CONFECCOES MATRIZ" in ws["A2"].value
    assert "01/2026" in ws["A2"].value

    # Linha 4 preenchida com os dados da NF 902
    assert ws["A4"].value == 1            # Sequencial
    assert str(ws["D4"].value) == "902"   # N. Fiscal
    assert float(ws["E4"].value) == 3000.0 # V. Total (3000 no xml_nf902)
    assert float(ws["F4"].value) == 3000.0 # Base de Cálculo (3000 no xml_nf902)
    assert float(ws["H4"].value) == 20.5   # A. DST (20.5%)
    assert float(ws["I4"].value) == 12.0   # A. ORI (12%)

    # Fórmulas de cálculo intactas na linha 4
    assert ws["J4"].value == "=(E4)/100*H4"
    assert ws["K4"].value == "=(F4)/100*I4"
    assert ws["L4"].value == "=J4-K4"

    # Colunas O (VR.PAGO) e P (DATA VENC.) vazias para preenchimento manual posterior
    assert ws["O4"].value is None
    assert ws["P4"].value is None

    # Totais no rodapé (Linha 123) intactos
    assert ws["E123"].value == "=SUM(E4:E122)"
    assert ws["F123"].value == "=SUM(F4:F122)"
    assert ws["G123"].value == "=SUM(G4:G122)"
    assert ws["J123"].value == "=SUM(J4:J122)"
    assert ws["K123"].value == "=SUM(K4:K122)"
    assert ws["L123"].value == "=SUM(L4:L122)"

    wb.close()
