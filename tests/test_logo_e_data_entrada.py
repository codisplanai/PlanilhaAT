import io
import os
import datetime
from decimal import Decimal
import openpyxl
import pytest

from app.core.config import settings
from app.services.excel.template_filler import TemplateFiller
from app.services.extraction.nfe_xml_extractor import NFeXMLExtractor
from tests.conftest import build_xml_nfe, CHAVE_NF901


def test_xml_extractor_nao_deduz_data_entrada_de_dhsaient():
    xml_com_saida = f"""<?xml version="1.0" encoding="UTF-8"?>
<nfeProc xmlns="http://www.portalfiscal.inf.br/nfe">
  <NFe>
    <infNFe Id="NFe{CHAVE_NF901}">
      <ide>
        <nNF>901</nNF>
        <serie>1</serie>
        <dhEmi>2026-06-10T10:00:00-03:00</dhEmi>
        <dhSaiEnt>2026-06-18T14:30:00-03:00</dhSaiEnt>
      </ide>
      <emit><CNPJ>98765432000180</CNPJ><enderEmit><UF>SP</UF></enderEmit></emit>
      <dest><CNPJ>12345678000195</CNPJ><enderDest><UF>BA</UF></enderDest></dest>
      <total><ICMSTot><vNF>1000.00</vNF><vBC>1000.00</vBC></ICMSTot></total>
      <det nItem="1">
        <prod><NCM>21069090</NCM><CFOP>6102</CFOP><vProd>1000.00</vProd></prod>
        <imposto><ICMS><ICMS00><vBC>1000.00</vBC><pICMS>12.00</pICMS></ICMS00></ICMS></imposto>
      </det>
    </infNFe>
  </NFe>
</nfeProc>""".encode("utf-8")

    extractor = NFeXMLExtractor()
    nf_data = extractor.extract_from_xml(xml_com_saida)

    # dhSaiEnt no XML é apenas saída do emitente; data_entrada da empresa deve ser None
    assert nf_data.data_entrada is None
    assert nf_data.data_emissao.date() == datetime.date(2026, 6, 10)


def test_template_filler_preserva_logo_em_qualquer_mes():
    template_path = os.path.join(settings.TEMPLATES_DIR, "modelo_padrao_antecipacao_parcial.xlsx")
    if not os.path.exists(template_path):
        pytest.skip("Template padrão não encontrado")

    mapping = {
        "start_row": 4,
        "columns": {
            "item_index": "A",
            "data_entrada": "B",
            "data_emissao": "C",
            "numero_nota": "D",
            "v_total": "E",
        },
        "header_cell": "A2",
    }

    for mes in ["01/2026", "06/2026", "08/2026", "12/2026"]:
        rows = [
            {
                "data_entrada": datetime.date(2026, int(mes[:2]), 10),
                "data_emissao": datetime.date(2026, int(mes[:2]), 5),
                "numero_nota": "999",
                "v_total": 500.0,
            }
        ]
        out_path = f"test_out_{mes.replace('/', '_')}.xlsx"
        try:
            TemplateFiller.fill_template(
                template_path=template_path,
                mapping=mapping,
                rows_data=rows,
                output_path=out_path,
                header_info={"razao_social": "EMPRESA TESTE LOGO", "competencia": mes},
            )

            wb = openpyxl.load_workbook(out_path)
            ws = wb.active
            # Verifica que a imagem/logo está presente na planilha gerada
            assert len(getattr(ws, "_images", [])) >= 1, f"Logo ausente para mês {mes}"

            # Verifica que a data de entrada foi inserida na coluna B
            val_b4 = ws["B4"].value
            if isinstance(val_b4, (datetime.datetime, datetime.date)):
                assert val_b4.strftime("%d/%m/%Y") == f"10/{mes[:2]}/2026"
            else:
                assert val_b4 == f"10/{mes[:2]}/2026"

            # Verifica que a data de emissão foi inserida na coluna C
            val_c4 = ws["C4"].value
            if isinstance(val_c4, (datetime.datetime, datetime.date)):
                assert val_c4.strftime("%d/%m/%Y") == f"05/{mes[:2]}/2026"

            wb.close()
        finally:
            if os.path.exists(out_path):
                os.remove(out_path)
