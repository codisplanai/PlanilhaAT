import os
import openpyxl
from decimal import Decimal
from datetime import date
import pytest
from fastapi.testclient import TestClient

from app.main import app
from app.core.database import Base, engine, SessionLocal
from app.models.perfil_regras import PerfilRegras
from app.models.empresa import Empresa
from app.models.template_xlsx import TemplateXlsx
from app.services.extraction.sped_fiscal_extractor import SpedFiscalExtractor
from app.services.excel.template_filler import TemplateFiller

client = TestClient(app)

@pytest.fixture
def db_session():
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()

def test_sped_extract_empresa_info_with_ie():
    """Testa extração de Razão Social, CNPJ, UF e Inscrição Estadual do Registro 0000 do SPED"""
    sped_sample = (
        "|0000|018|0|01042026|30042026|MERCADAO CARMO LTDA|12345678000195||MG|83592715|3106200|||A|0|\n"
        "|0001|0|\n"
        "|0990|3|\n"
    ).encode("utf-8")

    info = SpedFiscalExtractor.extract_empresa_info(sped_sample)
    assert info["razao_social"] == "MERCADAO CARMO LTDA"
    assert info["cnpj"] == "12345678000195"
    assert info["uf"] == "MG"
    assert info["ie"] == "83592715"

def test_empresa_api_crud_with_inscricao_estadual(db_session):
    """Testa criação, listagem e atualização de empresa com campo de Inscrição Estadual"""
    perfil = db_session.query(PerfilRegras).first()
    if not perfil:
        perfil = PerfilRegras(nome="Perfil Teste IE", padrao_uf={"SP": 0.12}, excecoes_ncm={})
        db_session.add(perfil)
        db_session.commit()
        db_session.refresh(perfil)

    # 1. Criação com IE
    cnpj_novo = "04252011000110"
    empresa_existente = db_session.query(Empresa).filter(Empresa.cnpj == cnpj_novo).first()
    if empresa_existente:
        db_session.delete(empresa_existente)
        db_session.commit()

    resp = client.post("/api/v1/empresas", json={
        "razao_social": "MERCADÃO CARMO",
        "cnpj": cnpj_novo,
        "inscricao_estadual": "83592715",
        "uf": "MG",
        "perfil_regras_id": perfil.id,
        "ativo": True
    })
    assert resp.status_code == 201
    created = resp.json()
    assert created["inscricao_estadual"] == "83592715"
    empresa_id = created["id"]

    # 2. Obtenção
    resp_get = client.get(f"/api/v1/empresas/{empresa_id}")
    assert resp_get.status_code == 200
    assert resp_get.json()["inscricao_estadual"] == "83592715"

    # 3. Atualização de IE
    resp_put = client.put(f"/api/v1/empresas/{empresa_id}", json={
        "inscricao_estadual": "99988877"
    })
    assert resp_put.status_code == 200
    assert resp_put.json()["inscricao_estadual"] == "99988877"

    # Limpeza
    client.delete(f"/api/v1/empresas/{empresa_id}")

def test_template_filler_header_formatting(tmp_path):
    """Testa preenchimento de cabeçalho (Empresa, IE, Competência) em arquivo Excel"""
    # Criar um workbook temporário com layout compatível
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "04-2026"
    ws["A1"] = "CONTABILIDADE"
    ws["A2"] = "Empresa -                                  IE:                  COMP."
    ws["B3"] = "DATA DE ENTRADA"
    ws["C3"] = "DATA DE EMISSÃO"
    ws["D3"] = "N. FISCAL"
    ws["E3"] = "V.TOTAL"

    template_file = os.path.join(tmp_path, "test_template.xlsx")
    wb.save(template_file)
    wb.close()

    output_file = os.path.join(tmp_path, "output_filled.xlsx")

    header_info = {
        "razao_social": "MERCADÃO CARMO",
        "ie": "83592715",
        "competencia": "04/2026",
        "mes": 4,
        "ano": 2026
    }

    mapping = {
        "start_row": 4,
        "header_cell": "A2",
        "columns": {
            "v_total": "E"
        }
    }

    rows_data = [
        {"v_total": Decimal("1500.00"), "data_emissao": date(2026, 4, 15)}
    ]

    TemplateFiller.fill_template(
        template_path=template_file,
        mapping=mapping,
        rows_data=rows_data,
        output_path=output_file,
        header_info=header_info
    )

    wb_out = openpyxl.load_workbook(output_file, data_only=False)
    assert len(wb_out.sheetnames) == 1
    assert wb_out.sheetnames == ["04-2026"]
    ws_out = wb_out["04-2026"]
    header_val = ws_out["A2"].value

    assert "MERCADÃO CARMO" in header_val
    assert "83592715" in header_val
    assert "04/2026" in header_val
    wb_out.close()

def test_template_filler_keeps_only_processed_sheet(tmp_path):
    """Testa que um template com 12 abas de meses gera um arquivo final com exatamente 1 aba da competência processada"""
    wb = openpyxl.Workbook()
    # Criar abas para todos os meses
    meses = ["JANEIRO", "FEVEREIRO", "MARÇO", "ABRIL", "MAIO", "JUNHO", "JULHO", "AGOSTO", "SETEMBRO", "OUTUBRO", "NOVEMBRO", "DEZEMBRO"]
    ws_first = wb.active
    ws_first.title = meses[0]
    for m in meses[1:]:
        wb.create_sheet(title=m)

    for m in meses:
        ws_m = wb[m]
        ws_m["A2"] = f"Empresa:                  Comp: {m}/2026"
        ws_m["B3"] = "DATA"
        ws_m["C3"] = "VALOR"
        ws_m["C4"] = "=C2+10" # fórmula para testar que não quebra

    template_multi = os.path.join(tmp_path, "template_12_meses.xlsx")
    wb.save(template_multi)
    wb.close()

    output_single = os.path.join(tmp_path, "output_abril_only.xlsx")

    header_info = {
        "razao_social": "MERCADÃO CARMO",
        "ie": "83592715",
        "competencia": "04/2026",
        "mes": 4,
        "ano": 2026
    }

    mapping = {
        "start_row": 5,
        "header_cell": "A2",
        "columns": {
            "v_total": "C"
        }
    }

    rows_data = [
        {"v_total": Decimal("250.00"), "data_emissao": date(2026, 4, 10)}
    ]

    TemplateFiller.fill_template(
        template_path=template_multi,
        mapping=mapping,
        rows_data=rows_data,
        output_path=output_single,
        header_info=header_info
    )

    # Verificar que a planilha resultante contém ESTRITAMENTE 1 aba ("ABRIL")
    wb_res = openpyxl.load_workbook(output_single, data_only=False)
    assert len(wb_res.sheetnames) == 1
    assert wb_res.sheetnames == ["ABRIL"]
    assert "MERCADÃO CARMO" in wb_res["ABRIL"]["A2"].value
    assert wb_res["ABRIL"]["C4"].value == "=C2+10"
    assert wb_res["ABRIL"]["C5"].value == 250.00
    wb_res.close()
