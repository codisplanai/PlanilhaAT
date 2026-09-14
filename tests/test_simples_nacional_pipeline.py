from decimal import Decimal
import os
import openpyxl
import pytest

from app.constants import (
    ANTECIPACAO_PARCIAL,
    ANTECIPACAO_PARCIAL_ANTECIPADO,
    ANTECIPACAO_PARCIAL_SIMPLES,
    ANTECIPACAO_PARCIAL_ANTECIPADO_SIMPLES,
)
from app.services.calculation.antecipacao_parcial import AntecipacaoParcialCalculator
from app.services.calculation.factory import CalculatorFactory
from app.services.excel.template_filler import TemplateFiller
from app.core.seeds import (
    DEFAULT_ANTECIPACAO_PARCIAL_SIMPLES_MAPPING,
    DEFAULT_ANTECIPACAO_PARCIAL_SIMPLES_TEMPLATE_PATH,
    DEFAULT_ANTECIPACAO_PARCIAL_ANTECIPADO_SIMPLES_MAPPING,
    DEFAULT_ANTECIPACAO_PARCIAL_ANTECIPADO_SIMPLES_TEMPLATE_PATH,
)


def test_antecipacao_parcial_calculator_simples_nacional_desconto_20():
    calc = AntecipacaoParcialCalculator()

    # Cenário Normal: V.Total = 1000.00, A.DST = 18%, A.ORI = 12%
    # Débito = 180.00, Crédito = 120.00, Valor Devido = 60.00
    res_normal = calc.calculate(
        v_total=Decimal("1000.00"),
        base_calculo=Decimal("1000.00"),
        ipi_despesas=Decimal("0.00"),
        a_ori=Decimal("0.12"),
        a_dst=Decimal("0.18"),
        parametros_extras={"is_simples": False},
    )
    assert res_normal.debito == Decimal("180.00")
    assert res_normal.credito == Decimal("120.00")
    assert res_normal.valor_devido == Decimal("60.00")
    assert res_normal.detalhes.get("reducao") == "0%"

    # Cenário Simples Nacional: 20% de desconto sobre 60.00 -> 48.00
    res_simples = calc.calculate(
        v_total=Decimal("1000.00"),
        base_calculo=Decimal("1000.00"),
        ipi_despesas=Decimal("0.00"),
        a_ori=Decimal("0.12"),
        a_dst=Decimal("0.18"),
        parametros_extras={"is_simples": True},
    )
    assert res_simples.debito == Decimal("180.00")
    assert res_simples.credito == Decimal("120.00")
    assert res_simples.valor_devido == Decimal("48.00")
    assert res_simples.detalhes.get("is_simples") is True
    assert res_simples.detalhes.get("reducao") == "20%"


def test_calculator_factory_returns_same_parcial_instance():
    calc_parcial = CalculatorFactory.get_calculator(ANTECIPACAO_PARCIAL)
    calc_simples = CalculatorFactory.get_calculator(ANTECIPACAO_PARCIAL_SIMPLES)
    calc_antecipado_simples = CalculatorFactory.get_calculator(ANTECIPACAO_PARCIAL_ANTECIPADO_SIMPLES)

    assert isinstance(calc_simples, AntecipacaoParcialCalculator)
    assert isinstance(calc_antecipado_simples, AntecipacaoParcialCalculator)
    assert calc_simples is calc_parcial


def test_fill_template_simples_preserva_formulas_e_desconto(tmp_path):
    assert os.path.exists(DEFAULT_ANTECIPACAO_PARCIAL_SIMPLES_TEMPLATE_PATH)

    output_path = os.path.join(tmp_path, "output_simples_test.xlsx")
    rows_data = [
        {
            "item_index": 1,
            "data_entrada": "01/02/2026",
            "data_emissao": "28/01/2026",
            "numero_nota": "12345",
            "v_total": Decimal("1000.00"),
            "base_calculo": Decimal("1000.00"),
            "ipi_despesas": Decimal("0.00"),
            "a_dst": Decimal("18.00"),
            "a_ori": Decimal("12.00"),
        }
    ]

    header_info = {
        "razao_social": "EMPRESA SIMPLES TESTE LTDA",
        "inscricao_estadual": "12345678",
        "competencia": "02/2026",
    }

    TemplateFiller.fill_template(
        template_path=DEFAULT_ANTECIPACAO_PARCIAL_SIMPLES_TEMPLATE_PATH,
        mapping=DEFAULT_ANTECIPACAO_PARCIAL_SIMPLES_MAPPING,
        rows_data=rows_data,
        output_path=output_path,
        header_info=header_info,
    )

    assert os.path.exists(output_path)
    wb = openpyxl.load_workbook(output_path, data_only=False)
    ws = wb.active

    # Verifica se as fórmulas de redução foram preservadas
    assert ws["J4"].value == "=(E4)/100*H4"
    assert ws["K4"].value == "=(F4)/100*I4"
    assert ws["L4"].value == "=J4-K4"
    assert ws["M4"].value == "=L4-(L4*$P$2)"
    assert ws["P2"].value == 0.2
    wb.close()


def test_fill_template_antecipado_simples_preserva_formulas(tmp_path):
    assert os.path.exists(DEFAULT_ANTECIPACAO_PARCIAL_ANTECIPADO_SIMPLES_TEMPLATE_PATH)

    output_path = os.path.join(tmp_path, "output_antecipado_simples_test.xlsx")
    rows_data = [
        {
            "item_index": 1,
            "data_entrada": None,
            "data_emissao": "15/02/2026",
            "numero_nota": "54321",
            "v_total": Decimal("2500.00"),
            "base_calculo": Decimal("2500.00"),
            "ipi_despesas": Decimal("0.00"),
            "a_dst": Decimal("20.50"),
            "a_ori": Decimal("12.00"),
        }
    ]

    header_info = {
        "razao_social": "EMPRESA SIMPLES TESTE LTDA",
        "inscricao_estadual": "12345678",
        "competencia": "02/2026",
    }

    TemplateFiller.fill_template(
        template_path=DEFAULT_ANTECIPACAO_PARCIAL_ANTECIPADO_SIMPLES_TEMPLATE_PATH,
        mapping=DEFAULT_ANTECIPACAO_PARCIAL_ANTECIPADO_SIMPLES_MAPPING,
        rows_data=rows_data,
        output_path=output_path,
        header_info=header_info,
    )

    assert os.path.exists(output_path)
    wb = openpyxl.load_workbook(output_path, data_only=False)
    ws = wb.active

    assert ws["B4"].value is None
    assert ws["D4"].value == "54321"
    assert ws["J4"].value == "=(E4)/100*H4"
    assert ws["K4"].value == "=(F4)/100*I4"
    assert ws["L4"].value == "=J4-K4"
    assert ws["M4"].value == "=L4-(L4*$P$2)"
    assert ws["P2"].value == 0.2
    wb.close()


def test_pipeline_routes_simples_nacional_and_applies_20_percent(
    db_session, cenario_janeiro, monkeypatch
):
    from datetime import datetime
    from app.models.empresa import Empresa
    from app.models.nota_fiscal import NotaFiscalProcessada
    from app.models.regra_cfop import RegraCfopDestino
    from app.services.extraction.base import ExtractedItemNF, ExtractedNFData
    from app.services.pipeline_service import ProcessingPipelineService

    # Cria cenário com os templates do Simples Nacional ativos
    solicitacao = cenario_janeiro(tipos=("antecipacao_parcial_simples",))
    empresa = db_session.query(Empresa).filter_by(id=solicitacao.empresa_id).one()
    # Ativa Simples Nacional na empresa
    empresa.optante_simples_nacional = True
    db_session.commit()

    perfil_id = empresa.perfil_regras_id
    db_session.add(
        RegraCfopDestino(perfil_regras_id=perfil_id, cfop_sufixo="102", destino="antecipacao_parcial")
    )
    db_session.commit()

    # Nota fiscal com CFOP 6102
    item = ExtractedItemNF(
        item_numero=1,
        ncm="84713012",
        cfop="6102",
        descricao="Produto Simples",
        descricao_confiavel=True,
        v_item=Decimal("1000.00"),
        v_total=Decimal("1000.00"),
        base_calculo=Decimal("1000.00"),
        ipi_despesas=Decimal("0.00"),
        a_ori=Decimal("0.12"),
    )
    nf = ExtractedNFData(
        numero_nota="777",
        serie="1",
        chave_acesso="0" * 44,
        cnpj_emitente="98765432000180",
        uf_emitente="SP",
        cnpj_destinatario=empresa.cnpj,
        uf_destinatario=empresa.uf,
        data_emissao=datetime(2026, 1, 15),
        v_total_nota=Decimal("1000.00"),
        v_bc_nota=Decimal("1000.00"),
        itens=[item],
    )

    pipeline = ProcessingPipelineService(db_session)
    monkeypatch.setattr(
        pipeline.source_loader,
        "load",
        lambda **kw: type("Sources", (), {
            "notes": [("dummy.xml", nf)],
            "entry_records": {},
            "sped_company_info": None,
            "ignored_notes": [],
        })(),
    )

    pipeline.process_solicitacao(
        solicitacao.id, xml_files_bytes=[("dummy.xml", b"<xml></xml>")]
    )

    db_session.refresh(solicitacao)
    notas = (
        db_session.query(NotaFiscalProcessada)
        .filter_by(solicitacao_id=solicitacao.id)
        .all()
    )
    assert len(notas) == 1
    n = notas[0]

    # Roteou para a planilha do Simples Nacional!
    assert n.destino_planilha == ANTECIPACAO_PARCIAL_SIMPLES
    # Débito = 1000 * 18% = 180.00, Crédito = 1000 * 12% = 120.00
    # Imposto bruto = 60.00 -> Com redução de 20% do Simples Nacional = 48.00!
    assert n.debito == Decimal("180.00")
    assert n.credito == Decimal("120.00")
    assert n.valor_devido == Decimal("48.00")

