import os
import json
import tempfile
import openpyxl
from decimal import Decimal
from datetime import datetime, date
import pytest
from sqlalchemy.orm import Session

from app.core.config import settings
from app.core.seeds import DEFAULT_ANTECIPACAO_TRIBUTARIA_MAPPING
from app.services.excel.template_filler import TemplateFiller
from app.services.excel.formula_guard import FormulaGuard
from app.services.rules_engine.mva_resolver import MvaResolver
from app.services.templates_admin.template_manager import TemplateManager
from app.services.pipeline_service import ProcessingPipelineService
from app.models.empresa import Empresa
from app.models.perfil_regras import PerfilRegras
from app.models.regra_aliquota import RegraAliquotaDestino
from app.models.solicitacao import Solicitacao

TEMPLATE_PATH = os.path.join(settings.TEMPLATES_DIR, "modelo_padrao_antecipacao_tributaria.xlsx")

def test_official_template_file_exists_and_has_sheets():
    assert os.path.exists(TEMPLATE_PATH), f"Arquivo do template {TEMPLATE_PATH} não encontrado"
    wb = openpyxl.load_workbook(TEMPLATE_PATH, data_only=False)
    assert len(wb.sheetnames) >= 1
    assert "04-2026" in wb.sheetnames or "01 (2)" in wb.sheetnames

    ws = wb["04-2026"] if "04-2026" in wb.sheetnames else wb.active
    assert "CÁLCULO ANTECIPAÇÃO TRIBUTÁRIA" in str(ws["F1"].value or ws["A1"].value)
    assert ws["B3"].value == "DATA DE ENTRADA"
    assert ws["H3"].value == "MVA"
    assert ws["N3"].value == "VALOR"
    assert ws["O3"].value == "DÉBITO"
    assert ws["P3"].value == "CRÉDITO"
    assert ws["Q3"].value == "VR.DEVIDO"

def test_template_formulas_intact_and_formula_guard():
    wb = openpyxl.load_workbook(TEMPLATE_PATH, data_only=False)
    formula_map = FormulaGuard.extract_formula_map(wb)
    assert len(formula_map) >= 1

    ws_target = wb["04-2026"] if "04-2026" in wb.sheetnames else wb.active
    sheet_formulas = formula_map[ws_target.title]
    assert len(sheet_formulas) >= 100, f"Aba {ws_target.title} deveria ter fórmulas mapeadas"
    
    # Verificar fórmulas de linhas de dados no modelo RP-151
    assert "IF(" in sheet_formulas[(4, 14)]                    # VALOR (Col N)
    assert sheet_formulas[(4, 15)] == "=N4*J4%"                 # DÉBITO (Col O)
    assert sheet_formulas[(4, 16)] == "=R4*K4%"                 # CRÉDITO (Col P)
    assert sheet_formulas[(4, 17)] == "=O4-P4"                  # VR.DEVIDO (Col Q)
    assert sheet_formulas[(4, 18)] == "=E4-G4"                  # Base sem IPI (Col R)
    assert sheet_formulas[(4, 20)] == "=Q4-S4"                  # DIFER (Col T)
    # Verificar totais
    assert sheet_formulas[(39, 5)] == "=SUM(E4:E38)"            # TOTAL V.TOTAL
    assert sheet_formulas[(39, 14)] == "=SUM(N4:N38)"          # TOTAL VALOR
    assert sheet_formulas[(39, 17)] == "=SUM(Q4:Q38)"          # TOTAL VR.DEVIDO

def test_mva_resolver_behavior(tmp_path):
    # 1. Sem arquivo de Anexo I (fallback para 0.0)
    non_existent = str(tmp_path / "NonExistentAnexo.json")
    MvaResolver.set_custom_path(non_existent)
    assert MvaResolver.resolve_mva("84189900") == Decimal("0.00")
    assert MvaResolver.resolve_mva("84189900", fallback_mva=Decimal("50.0")) == Decimal("50.0")

    # 2. Com arquivo JSON estruturado (dict)
    sample_json = str(tmp_path / "AnexoI.json")
    with open(sample_json, "w", encoding="utf-8") as f:
        json.dump({
            "84189900": 57.92,
            "84159090": 69.62,
            "8414": 45.00
        }, f)

    MvaResolver.set_custom_path(sample_json)
    assert MvaResolver.resolve_mva("8418.99.00") == Decimal("57.92")
    assert MvaResolver.resolve_mva("84159090") == Decimal("69.62")
    # Match por prefixo de 4 dígitos
    assert MvaResolver.resolve_mva("84145990") == Decimal("45.00")
    # NCM não encontrado
    assert MvaResolver.resolve_mva("99999999") == Decimal("0.00")

    # 3. Teste com o arquivo real AnexoI.json (storage/data/AnexoI.json)
    real_anexo_path = os.path.join(settings.STORAGE_DIR, "data", "AnexoI.json")
    if os.path.exists(real_anexo_path):
        MvaResolver.set_custom_path(real_anexo_path)
        # NCM 7308901 (barras para construção)
        assert MvaResolver.resolve_mva("7308.90.10", a_ori=Decimal("0.12")) == Decimal("49.43")
        assert MvaResolver.resolve_mva("73089010", a_ori=Decimal("0.07")) == Decimal("57.92")
        assert MvaResolver.resolve_mva("73089010", a_ori=Decimal("0.04")) == Decimal("63.02")

        # NCM 73083 (portas e janelas)
        assert MvaResolver.resolve_mva("7308.30.00", a_ori=Decimal("0.12")) == Decimal("60.50")
        assert MvaResolver.resolve_mva("73083000", a_ori=Decimal("0.07")) == Decimal("69.62")
        assert MvaResolver.resolve_mva("73083000", a_ori=Decimal("0.04")) == Decimal("75.09")

        # NCM 22011 (água mineral)
        assert MvaResolver.resolve_mva("2201.10.00", a_ori=Decimal("12.0")) == Decimal("136.88")
        assert MvaResolver.resolve_mva("22011000", a_ori=Decimal("7.0")) == Decimal("150.34")
        assert MvaResolver.resolve_mva("22011000", a_ori=Decimal("4.0")) == Decimal("158.42")

def test_template_filler_populates_antecipacao_tributaria_cleanly(tmp_path):
    output_xlsx = str(tmp_path / "test_out_tributaria.xlsx")
    rows_data = [
        {
            "numero_nota": "77084",
            "data_entrada": date(2026, 5, 20),
            "data_emissao": date(2026, 5, 16),
            "v_total": Decimal("24848.13"),
            "base_calculo": Decimal("24848.13"),
            "ipi_despesas": Decimal("0.00"),
            "mva": Decimal("49.43"),
            "reducao": None,
            "a_dst": Decimal("20.50"),
            "a_ori": Decimal("12.00"),
            "red": "",
            "debito": Decimal("7611.76"),
            "credito": Decimal("2981.78"),
            "valor_devido": Decimal("4629.98")
        },
        {
            "numero_nota": "113251",
            "data_entrada": date(2026, 5, 22),
            "data_emissao": date(2026, 5, 18),
            "v_total": Decimal("4564.33"),
            "base_calculo": Decimal("4446.97"),
            "ipi_despesas": Decimal("117.36"),
            "mva": Decimal("75.09"),
            "reducao": None,
            "a_dst": Decimal("20.50"),
            "a_ori": Decimal("4.00"),
            "red": "",
            "debito": Decimal("1638.30"),
            "credito": Decimal("177.88"),
            "valor_devido": Decimal("1460.42")
        }
    ]

    header_info = {
        "razao_social": "ELETRONE REFRIGERACAO E PECAS LTDA",
        "cnpj": "08089915000100",
        "ie": "08089915",
        "competencia": "04/2026",
        "mes": 4,
        "ano": 2026
    }

    TemplateFiller.fill_template(
        template_path=TEMPLATE_PATH,
        mapping=DEFAULT_ANTECIPACAO_TRIBUTARIA_MAPPING,
        rows_data=rows_data,
        output_path=output_xlsx,
        header_info=header_info
    )

    assert os.path.exists(output_xlsx)
    out_wb = openpyxl.load_workbook(output_xlsx, data_only=False)
    assert len(out_wb.sheetnames) == 1
    assert "04-2026" in out_wb.sheetnames
    ws_out = out_wb["04-2026"]

    # Validar header
    assert "ELETRONE REFRIGERACAO" in ws_out["A2"].value
    assert "04/2026" in ws_out["A2"].value

    # Linha 4
    assert ws_out["A4"].value == 1
    assert ws_out["D4"].value == "77084"
    assert ws_out["E4"].value == 24848.13
    assert ws_out["F4"].value == 24848.13
    assert ws_out["H4"].value == 49.43
    assert ws_out["J4"].value == 20.50
    assert ws_out["K4"].value == 12.00
    assert "IF(" in ws_out["N4"].value
    assert ws_out["O4"].value == "=N4*J4%"
    assert ws_out["P4"].value == "=R4*K4%"
    assert ws_out["Q4"].value == "=O4-P4"
    assert ws_out["R4"].value == "=E4-G4"

    # Linha 5
    assert ws_out["A5"].value == 2
    assert ws_out["D5"].value == "113251"
    assert ws_out["E5"].value == 4564.33
    assert ws_out["F5"].value == 4446.97
    assert ws_out["G5"].value == 117.36
    assert ws_out["H5"].value == 75.09
    assert "IF(" in ws_out["N5"].value

    # Validar que a integridade total do workbook foi preservada
    orig_wb = openpyxl.load_workbook(TEMPLATE_PATH, data_only=False)
    orig_fmap = FormulaGuard.extract_formula_map(orig_wb)
    FormulaGuard.verify_wb_integrity(orig_fmap, out_wb)
