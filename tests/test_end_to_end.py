import io
import json
import openpyxl
from decimal import Decimal
from datetime import datetime, date

def test_fluxo_completo_api_end_to_end(client, sample_xml_nfe, create_sample_excel_template):
    # 1. Cadastrar Perfil de Regras
    res_perfil = client.post("/api/v1/perfis-regras", json={
        "nome": "Perfil Comércio Geral BA",
        "descricao": "Perfil para empresas varejistas",
        "configuracoes_extras": {"fração_base": 1.0}
    })
    assert res_perfil.status_code == 201, res_perfil.text
    perfil_id = res_perfil.json()["id"]

    # 2. Cadastrar Regras de Alíquotas de Destino (A.DST)
    # Regra padrão estadual BA (18%)
    res_regra_padrao = client.post("/api/v1/regras-aliquotas", json={
        "perfil_regras_id": perfil_id,
        "uf": "BA",
        "ncm": None,
        "aliquota": 0.1800,
        "descricao": "Alíquota padrão BA"
    })
    assert res_regra_padrao.status_code == 201, res_regra_padrao.text

    # Regra exceção para NCM 84713012 (20.5%)
    res_regra_excecao = client.post("/api/v1/regras-aliquotas", json={
        "perfil_regras_id": perfil_id,
        "uf": "BA",
        "ncm": "84713012",
        "aliquota": 0.2050,
        "descricao": "Exceção Informática BA"
    })
    assert res_regra_excecao.status_code == 201, res_regra_excecao.text

    # 3. Cadastrar Empresa (CNPJ: 12345678000195, que bate com o destinatário do sample_xml_nfe)
    res_empresa = client.post("/api/v1/empresas", json={
        "razao_social": "Cliente Contabilidade Bahia LTDA",
        "cnpj": "12.345.678/0001-95",
        "uf": "BA",
        "perfil_regras_id": perfil_id,
        "ativo": True
    })
    assert res_empresa.status_code == 201, res_empresa.text
    empresa_id = res_empresa.json()["id"]

    # 4. Upload de Template Excel para Antecipação Parcial
    template_file_path = create_sample_excel_template("antecipacao_parcial")
    with open(template_file_path, "rb") as f:
        template_bytes = f.read()

    mapping = {
        "start_row": 4,
        "sheet_name": "Planilha AT",
        "columns": {
            "numero_nota": "A",
            "data_emissao": "B",
            "ncm": "C",
            "v_total": "D",
            "base_calculo": "E",
            "a_ori": "F",
            "a_dst": "G"
        }
    }

    res_template = client.post(
        "/api/v1/templates/upload",
        data={
            "tipo": "antecipacao_parcial",
            "mapeamento_json": json.dumps(mapping),
            "observacoes": "Template v1 oficial",
            "promover_ativo": True
        },
        files={
            "file": ("template_parcial_v1.xlsx", template_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        }
    )
    assert res_template.status_code == 201, res_template.text
    template_id = res_template.json()["id"]

    # 5. Criar Solicitação de Processamento
    res_solicitacao = client.post("/api/v1/solicitacoes", json={
        "empresa_id": empresa_id,
        "periodo_inicio": "2026-01-01",
        "periodo_fim": "2026-01-31",
        "tipo_planilha": "antecipacao_parcial"
    })
    assert res_solicitacao.status_code == 201, res_solicitacao.text
    solicitacao_id = res_solicitacao.json()["id"]

    # 6. Upload dos XMLs e Processamento
    res_proc = client.post(
        f"/api/v1/solicitacoes/{solicitacao_id}/processar",
        files=[
            ("files", ("nfe_1234.xml", sample_xml_nfe, "text/xml"))
        ]
    )
    assert res_proc.status_code == 200, res_proc.text
    solicitacao_data = res_proc.json()
    assert solicitacao_data["status"] == "concluido"
    assert solicitacao_data["total_notas_processadas"] == 1

    # 7. Download da Planilha Gerada
    res_download = client.get(f"/api/v1/solicitacoes/{solicitacao_id}/download")
    assert res_download.status_code == 200
    assert len(res_download.content) > 0

    # 8. Inspeção do arquivo Excel baixado
    wb = openpyxl.load_workbook(io.BytesIO(res_download.content), data_only=False)
    ws = wb["Planilha AT"]

    # Dados inseridos na linha 4
    assert str(ws["A4"].value) == "1234"
    if isinstance(ws["B4"].value, (datetime, date)):
        assert ws["B4"].value.strftime("%d/%m/%Y") == "15/01/2026"
    else:
        assert ws["B4"].value == "15/01/2026"
    assert ws["C4"].value == "84713012"
    assert float(ws["D4"].value) == 5400.0
    assert float(ws["E4"].value) == 5150.0
    assert float(ws["F4"].value) == 0.12   # A.ORI extraída do XML (pICMS 12%)
    assert float(ws["G4"].value) == 0.205  # A.DST resolvida pelo motor (exceção NCM 84713012)

    # Fórmulas intactas
    assert ws["H4"].value == "=D4*G4"
    assert ws["I4"].value == "=E4*F4"
    assert ws["J4"].value == "=H4-I4"
    assert ws["D12"].value == "=SUM(D4:D10)"

    wb.close()


def test_fluxo_com_notas_ignoradas_fora_do_periodo(client, sample_xml_nfe, create_sample_excel_template):
    # 1. Setup básico de Empresa e Regras
    res_perfil = client.post("/api/v1/perfis-regras", json={
        "nome": "Perfil Teste Ignorar Fora Periodo",
        "configuracoes_extras": {}
    })
    perfil_id = res_perfil.json()["id"]

    client.post("/api/v1/regras-aliquotas", json={
        "perfil_regras_id": perfil_id,
        "uf": "BA",
        "ncm": None,
        "aliquota": 0.18,
        "descricao": "Padrão"
    })

    res_empresa = client.post("/api/v1/empresas", json={
        "razao_social": "Empresa Teste Periodo LTDA",
        "cnpj": "12.345.678/0001-95",
        "uf": "BA",
        "perfil_regras_id": perfil_id,
        "ativo": True
    })
    empresa_id = res_empresa.json()["id"]

    # 2. Upload template
    template_file_path = create_sample_excel_template("antecipacao_parcial")
    with open(template_file_path, "rb") as f:
        template_bytes = f.read()

    mapping = {
        "start_row": 4,
        "sheet_name": "Planilha AT",
        "columns": {
            "numero_nota": "A",
            "data_emissao": "B",
            "ncm": "C",
            "v_total": "D",
            "base_calculo": "E",
            "a_ori": "F",
            "a_dst": "G"
        }
    }

    client.post(
        "/api/v1/templates/upload",
        data={
            "tipo": "antecipacao_parcial",
            "mapeamento_json": json.dumps(mapping),
            "observacoes": "Template v1",
            "promover_ativo": True
        },
        files={
            "file": ("template_parcial_v1.xlsx", template_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        }
    )

    # 3. XML 1: Válido (emissão 15/01/2026)
    xml_valido = sample_xml_nfe

    # 4. XML 2: Fora do período (emissão 19/05/2026)
    xml_fora = sample_xml_nfe.decode("utf-8").replace("2026-01-15", "2026-05-19").replace("<nNF>1234</nNF>", "<nNF>104112</nNF>").encode("utf-8")

    # 5. Criar solicitação para Janeiro (2026-01-01 a 2026-01-31)
    res_solicitacao = client.post("/api/v1/solicitacoes", json={
        "empresa_id": empresa_id,
        "periodo_inicio": "2026-01-01",
        "periodo_fim": "2026-01-31",
        "tipo_planilha": "antecipacao_parcial"
    })
    solicitacao_id = res_solicitacao.json()["id"]

    # 6. Processar ambos os XMLs juntos
    res_proc = client.post(
        f"/api/v1/solicitacoes/{solicitacao_id}/processar",
        files=[
            ("files", ("nfe_1234.xml", xml_valido, "text/xml")),
            ("files", ("nfe_104112.xml", xml_fora, "text/xml")),
        ]
    )
    assert res_proc.status_code == 200, res_proc.text
    dados = res_proc.json()

    # Deve concluir com sucesso, tendo 1 nota processada e 1 nota ignorada
    assert dados["status"] == "concluido"
    assert dados["total_notas_processadas"] == 1
    assert len(dados["notas_ignoradas"]) == 1
    assert dados["notas_ignoradas"][0]["numero_nota"] == "104112"
    assert "fora do período informado" in dados["notas_ignoradas"][0]["motivo"]

    # 7. Caso extremo: se todos os XMLs forem fora do período, deve retornar erro explicativo 400
    res_solic_fevereiro = client.post("/api/v1/solicitacoes", json={
        "empresa_id": empresa_id,
        "periodo_inicio": "2026-02-01",
        "periodo_fim": "2026-02-28",
        "tipo_planilha": "antecipacao_parcial"
    })
    solic_fev_id = res_solic_fevereiro.json()["id"]

    res_proc_todos_fora = client.post(
        f"/api/v1/solicitacoes/{solic_fev_id}/processar",
        files=[
            ("files", ("nfe_104112.xml", xml_fora, "text/xml")),
        ]
    )
    assert res_proc_todos_fora.status_code == 400
    assert "Todas as 1 nota(s) enviadas foram desconsideradas" in res_proc_todos_fora.json()["detail"]


def test_preenchimento_excel_ordem_crescente_data_emissao(client, sample_xml_nfe, create_sample_excel_template):
    # 1. Setup básico de Empresa e Regras
    res_perfil = client.post("/api/v1/perfis-regras", json={
        "nome": "Perfil Teste Ordenacao Emissao",
        "configuracoes_extras": {}
    })
    perfil_id = res_perfil.json()["id"]

    client.post("/api/v1/regras-aliquotas", json={
        "perfil_regras_id": perfil_id,
        "uf": "BA",
        "ncm": None,
        "aliquota": 0.18,
        "descricao": "Padrão"
    })

    res_empresa = client.post("/api/v1/empresas", json={
        "razao_social": "Empresa Teste Ordem Emissao LTDA",
        "cnpj": "12.345.678/0001-95",
        "uf": "BA",
        "perfil_regras_id": perfil_id,
        "ativo": True
    })
    empresa_id = res_empresa.json()["id"]

    # 2. Upload template
    template_file_path = create_sample_excel_template("antecipacao_parcial")
    with open(template_file_path, "rb") as f:
        template_bytes = f.read()

    mapping = {
        "start_row": 4,
        "sheet_name": "Planilha AT",
        "columns": {
            "numero_nota": "A",
            "data_emissao": "B",
            "ncm": "C",
            "v_total": "D",
            "base_calculo": "E",
            "a_ori": "F",
            "a_dst": "G"
        }
    }

    client.post(
        "/api/v1/templates/upload",
        data={
            "tipo": "antecipacao_parcial",
            "mapeamento_json": json.dumps(mapping),
            "observacoes": "Template v1",
            "promover_ativo": True
        },
        files={
            "file": ("template_parcial_v1.xlsx", template_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        }
    )

    # 3. Criar 3 XMLs com datas desordenadas:
    # NF 300 -> 25/01/2026
    xml_25 = sample_xml_nfe.decode("utf-8").replace("2026-01-15", "2026-01-25").replace("<nNF>1234</nNF>", "<nNF>300</nNF>").encode("utf-8")
    # NF 100 -> 05/01/2026
    xml_05 = sample_xml_nfe.decode("utf-8").replace("2026-01-15", "2026-01-05").replace("<nNF>1234</nNF>", "<nNF>100</nNF>").encode("utf-8")
    # NF 200 -> 18/01/2026
    xml_18 = sample_xml_nfe.decode("utf-8").replace("2026-01-15", "2026-01-18").replace("<nNF>1234</nNF>", "<nNF>200</nNF>").encode("utf-8")

    # 4. Criar solicitação
    res_solic = client.post("/api/v1/solicitacoes", json={
        "empresa_id": empresa_id,
        "periodo_inicio": "2026-01-01",
        "periodo_fim": "2026-01-31",
        "tipo_planilha": "antecipacao_parcial"
    })
    solic_id = res_solic.json()["id"]

    # 5. Enviar em ordem fora de sequência cronológica: [25/01, 05/01, 18/01]
    res_proc = client.post(
        f"/api/v1/solicitacoes/{solic_id}/processar",
        files=[
            ("files", ("nfe_300.xml", xml_25, "text/xml")),
            ("files", ("nfe_100.xml", xml_05, "text/xml")),
            ("files", ("nfe_200.xml", xml_18, "text/xml")),
        ]
    )
    assert res_proc.status_code == 200, res_proc.text
    dados = res_proc.json()
    assert dados["total_notas_processadas"] == 3

    # Verificar que na resposta da API as notas estão em ordem crescente de emissão
    notas_api = dados["notas_processadas"]
    assert [n["numero_nota"] for n in notas_api] == ["100", "200", "300"]

    # 6. Baixar a planilha Excel gerada e conferir a ordem das linhas escritas
    res_download = client.get(f"/api/v1/solicitacoes/{solic_id}/download")
    assert res_download.status_code == 200

    wb = openpyxl.load_workbook(io.BytesIO(res_download.content), data_only=False)
    ws = wb["Planilha AT"]

    # Linha 4 -> NF 100 (05/01/2026)
    assert str(ws["A4"].value) == "100"
    val_b4 = ws["B4"].value.strftime("%d/%m/%Y") if isinstance(ws["B4"].value, (datetime, date)) else ws["B4"].value
    assert val_b4 == "05/01/2026"

    # Linha 5 -> NF 200 (18/01/2026)
    assert str(ws["A5"].value) == "200"
    val_b5 = ws["B5"].value.strftime("%d/%m/%Y") if isinstance(ws["B5"].value, (datetime, date)) else ws["B5"].value
    assert val_b5 == "18/01/2026"

    # Linha 6 -> NF 300 (25/01/2026)
    assert str(ws["A6"].value) == "300"
    val_b6 = ws["B6"].value.strftime("%d/%m/%Y") if isinstance(ws["B6"].value, (datetime, date)) else ws["B6"].value
    assert val_b6 == "25/01/2026"

    wb.close()


def test_preenchimento_excel_ordem_crescente_data_entrada_preferencial(client, sample_xml_nfe, create_sample_excel_template):
    # 1. Setup básico
    res_perfil = client.post("/api/v1/perfis-regras", json={
        "nome": "Perfil Teste Data Entrada Preferencial",
        "configuracoes_extras": {}
    })
    perfil_id = res_perfil.json()["id"]

    client.post("/api/v1/regras-aliquotas", json={
        "perfil_regras_id": perfil_id,
        "uf": "BA",
        "ncm": None,
        "aliquota": 0.18,
        "descricao": "Padrão"
    })

    res_empresa = client.post("/api/v1/empresas", json={
        "razao_social": "Empresa Teste Data Entrada LTDA",
        "cnpj": "12.345.678/0001-95",
        "uf": "BA",
        "perfil_regras_id": perfil_id,
        "ativo": True
    })
    empresa_id = res_empresa.json()["id"]

    template_file_path = create_sample_excel_template("antecipacao_parcial")
    with open(template_file_path, "rb") as f:
        template_bytes = f.read()

    mapping = {
        "start_row": 4,
        "sheet_name": "Planilha AT",
        "columns": {
            "numero_nota": "A",
            "data_emissao": "B",
            "ncm": "C",
            "v_total": "D",
            "base_calculo": "E",
            "a_ori": "F",
            "a_dst": "G"
        }
    }

    client.post(
        "/api/v1/templates/upload",
        data={
            "tipo": "antecipacao_parcial",
            "mapeamento_json": json.dumps(mapping),
            "observacoes": "Template v1",
            "promover_ativo": True
        },
        files={
            "file": ("template_parcial_v1.xlsx", template_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        }
    )

    # 2. Criar 3 XMLs com emissão invertida em relação à entrada:
    # NF 100 -> Emissão: 05/01/2026, Entrada na planilha: 20/02/2026 (Será a 3ª linha)
    # NF 200 -> Emissão: 15/01/2026, Entrada na planilha: 10/02/2026 (Será a 2ª linha)
    # NF 300 -> Emissão: 25/01/2026, Entrada na planilha: 01/02/2026 (Será a 1ª linha)
    xml_100 = sample_xml_nfe.decode("utf-8").replace("2026-01-15", "2026-01-05").replace("<nNF>1234</nNF>", "<nNF>100</nNF>").encode("utf-8")
    xml_200 = sample_xml_nfe.decode("utf-8").replace("2026-01-15", "2026-01-15").replace("<nNF>1234</nNF>", "<nNF>200</nNF>").encode("utf-8")
    xml_300 = sample_xml_nfe.decode("utf-8").replace("2026-01-15", "2026-01-25").replace("<nNF>1234</nNF>", "<nNF>300</nNF>").encode("utf-8")

    # 3. Criar a planilha de datas de entrada do sistema contábil
    wb_ent = openpyxl.Workbook()
    ws_ent = wb_ent.active
    ws_ent.append(["CNPJ Emitente", "Série", "Número Nota", "Data Entrada"])
    ws_ent.append(["98765432000180", "1", "100", "20/02/2026"])
    ws_ent.append(["98765432000180", "1", "200", "10/02/2026"])
    ws_ent.append(["98765432000180", "1", "300", "01/02/2026"])
    buf = io.BytesIO()
    wb_ent.save(buf)
    planilha_entradas_bytes = buf.getvalue()

    # 4. Criar solicitação
    res_solic = client.post("/api/v1/solicitacoes", json={
        "empresa_id": empresa_id,
        "periodo_inicio": "2026-01-01",
        "periodo_fim": "2026-01-31",
        "tipo_planilha": "antecipacao_parcial"
    })
    solic_id = res_solic.json()["id"]

    # 5. Processar enviando XMLs e planilha de entrada contábil
    res_proc = client.post(
        f"/api/v1/solicitacoes/{solic_id}/processar",
        files=[
            ("files", ("nfe_100.xml", xml_100, "text/xml")),
            ("files", ("nfe_200.xml", xml_200, "text/xml")),
            ("files", ("nfe_300.xml", xml_300, "text/xml")),
            ("planilha_entradas", ("entradas_contabil.xlsx", planilha_entradas_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"))
        ]
    )
    assert res_proc.status_code == 200, res_proc.text
    dados = res_proc.json()

    # Como a ordem por Data de Entrada é:
    # 1º: NF 300 (01/02/2026)
    # 2º: NF 200 (10/02/2026)
    # 3º: NF 100 (20/02/2026)
    notas_api = dados["notas_processadas"]
    assert [n["numero_nota"] for n in notas_api] == ["300", "200", "100"]

    # 6. Conferir Excel gerado
    res_download = client.get(f"/api/v1/solicitacoes/{solic_id}/download")
    assert res_download.status_code == 200

    wb = openpyxl.load_workbook(io.BytesIO(res_download.content), data_only=False)
    ws = wb["Planilha AT"]

    assert str(ws["A4"].value) == "300"
    assert str(ws["A5"].value) == "200"
    assert str(ws["A6"].value) == "100"

    wb.close()



