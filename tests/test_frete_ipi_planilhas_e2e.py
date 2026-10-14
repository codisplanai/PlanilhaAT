import os
import openpyxl
from decimal import Decimal
from datetime import date

from app.models.perfil_regras import PerfilRegras
from app.models.empresa import Empresa
from app.models.regra_aliquota import RegraAliquotaDestino
from app.models.solicitacao import Solicitacao
from app.models.template_xlsx import TemplateXlsx
from app.core.seeds import seed_default_templates
from app.services.pipeline_service import ProcessingPipelineService

def test_e2e_antecipacao_tributaria_planilha_keeps_frete_out_of_ipi(db_session):
    """
    Na Antecipação Tributária, a coluna G chama-se 'IPI'.
    Se a nota tiver frete mas não tiver IPI destacado, a coluna G deve ser 0.00.
    """
    seed_default_templates(db_session)
    template = (
        db_session.query(TemplateXlsx)
        .filter(TemplateXlsx.tipo == "antecipacao_tributaria", TemplateXlsx.ativo == True)
        .first()
    )
    assert template is not None

    perfil = PerfilRegras(nome="Perfil AT E2E")
    db_session.add(perfil)
    db_session.commit()

    db_session.add(RegraAliquotaDestino(
        perfil_regras_id=perfil.id,
        uf="BA",
        aliquota=Decimal("0.2050"),
        descricao="Alíquota BA"
    ))
    empresa = Empresa(
        razao_social="TESTE AT LTDA",
        cnpj="12345678000195",
        uf="BA",
        perfil_regras_id=perfil.id
    )
    db_session.add(empresa)
    db_session.commit()

    # XML de fornecedor SP para BA: vProd 1000, Frete 100, Base 1100, vNF 1100, IPI 0.
    # CFOP 6403 ou NCM com ST para ir para antecipacao_tributaria
    xml = """<?xml version="1.0" encoding="UTF-8"?>
    <nfeProc xmlns="http://www.portalfiscal.inf.br/nfe">
      <NFe>
        <infNFe Id="NFe35260198765432000180550010000005551000000555">
          <ide><nNF>555</nNF><serie>1</serie><dhEmi>2026-01-10T10:00:00-03:00</dhEmi></ide>
          <emit><CNPJ>98765432000180</CNPJ><enderEmit><UF>SP</UF></enderEmit><CRT>3</CRT></emit>
          <dest><CNPJ>12345678000195</CNPJ><enderDest><UF>BA</UF></enderDest></dest>
          <det nItem="1">
            <prod>
              <NCM>21069090</NCM><CFOP>6403</CFOP><xProd>MERCADORIA ST</xProd>
              <vProd>1000.00</vProd><vFrete>100.00</vFrete><vSeg>0.00</vSeg><vOutro>0.00</vOutro><vDesc>0.00</vDesc>
            </prod>
            <imposto>
              <ICMS><ICMS10><orig>0</orig><CST>10</CST><vBC>1100.00</vBC><pICMS>12.00</pICMS><vICMS>132.00</vICMS></ICMS10></ICMS>
            </imposto>
          </det>
          <total><ICMSTot><vBC>1100.00</vBC><vICMS>132.00</vICMS><vProd>1000.00</vProd><vFrete>100.00</vFrete><vSeg>0.00</vSeg><vOutro>0.00</vOutro><vIPI>0.00</vIPI><vNF>1100.00</vNF></ICMSTot></total>
        </infNFe>
      </NFe>
    </nfeProc>""".encode("utf-8")

    solicitacao = Solicitacao(
        empresa_id=empresa.id,
        periodo_inicio=date(2026, 1, 1),
        periodo_fim=date(2026, 1, 31),
        tipo_planilha="antecipacao_tributaria",
        template_id=template.id,
        status="pendente"
    )
    db_session.add(solicitacao)
    db_session.commit()

    service = ProcessingPipelineService(db_session)
    res = service.process_solicitacao(solicitacao.id, xml_files_bytes=[("555.xml", xml)])
    assert res.status == "concluido"
    assert os.path.exists(res.arquivo_saida_path)

    wb = openpyxl.load_workbook(res.arquivo_saida_path, data_only=False)
    ws = wb.active
    # Coluna E = V.TOTAL (1100.00)
    assert float(ws["E4"].value) == 1100.0
    # Coluna F = B.CALCULO (1100.00)
    assert float(ws["F4"].value) == 1100.0
    # Coluna G = IPI (DEVE SER 0.00 ou None, NUNCA 100.00!)
    val_g4 = ws["G4"].value
    assert val_g4 is None or float(val_g4) == 0.0
    wb.close()

def test_e2e_difal_planilha_keeps_frete_out_of_ipi(db_session):
    """
    No DIFAL, a coluna F é IPI + DESP.
    Se a nota tiver frete mas não tiver IPI destacado, a coluna F deve ser 0.00,
    evitando que o frete seja duplicado no cálculo de Base ST da fórmula Excel.
    """
    seed_default_templates(db_session)
    template = (
        db_session.query(TemplateXlsx)
        .filter(TemplateXlsx.tipo == "difal", TemplateXlsx.ativo == True)
        .first()
    )
    assert template is not None

    perfil = PerfilRegras(nome="Perfil DIFAL E2E")
    db_session.add(perfil)
    db_session.commit()

    db_session.add(RegraAliquotaDestino(
        perfil_regras_id=perfil.id,
        uf="BA",
        aliquota=Decimal("0.2050"),
        descricao="Alíquota BA"
    ))
    empresa = Empresa(
        razao_social="TESTE DIFAL LTDA",
        cnpj="12345678000195",
        uf="BA",
        perfil_regras_id=perfil.id
    )
    db_session.add(empresa)
    db_session.commit()

    # CFOP 6556 / 6407 = Uso e Consumo -> vai para DIFAL
    xml = """<?xml version="1.0" encoding="UTF-8"?>
    <nfeProc xmlns="http://www.portalfiscal.inf.br/nfe">
      <NFe>
        <infNFe Id="NFe35260198765432000180550010000007771000000777">
          <ide><nNF>777</nNF><serie>1</serie><dhEmi>2026-01-10T10:00:00-03:00</dhEmi></ide>
          <emit><CNPJ>98765432000180</CNPJ><enderEmit><UF>SP</UF></enderEmit><CRT>3</CRT></emit>
          <dest><CNPJ>12345678000195</CNPJ><enderDest><UF>BA</UF></enderDest></dest>
          <det nItem="1">
            <prod>
              <NCM>84713012</NCM><CFOP>6556</CFOP><xProd>MATERIAL USO E CONSUMO</xProd>
              <vProd>1000.00</vProd><vFrete>100.00</vFrete><vSeg>0.00</vSeg><vOutro>0.00</vOutro><vDesc>0.00</vDesc>
            </prod>
            <imposto>
              <ICMS><ICMS00><orig>0</orig><CST>00</CST><vBC>1100.00</vBC><pICMS>12.00</pICMS><vICMS>132.00</vICMS></ICMS00></ICMS>
            </imposto>
          </det>
          <total><ICMSTot><vBC>1100.00</vBC><vICMS>132.00</vICMS><vProd>1000.00</vProd><vFrete>100.00</vFrete><vSeg>0.00</vSeg><vOutro>0.00</vOutro><vIPI>0.00</vIPI><vNF>1100.00</vNF></ICMSTot></total>
        </infNFe>
      </NFe>
    </nfeProc>""".encode("utf-8")

    solicitacao = Solicitacao(
        empresa_id=empresa.id,
        periodo_inicio=date(2026, 1, 1),
        periodo_fim=date(2026, 1, 31),
        tipo_planilha="difal",
        template_id=template.id,
        status="pendente"
    )
    db_session.add(solicitacao)
    db_session.commit()

    service = ProcessingPipelineService(db_session)
    res = service.process_solicitacao(solicitacao.id, xml_files_bytes=[("777.xml", xml)])
    assert res.status == "concluido"
    assert os.path.exists(res.arquivo_saida_path)

    wb = openpyxl.load_workbook(res.arquivo_saida_path, data_only=False)
    # DIFAL salva na aba do mês correspondente (JANEIRO)
    assert "JANEIRO" in wb.sheetnames
    ws = wb["JANEIRO"]
    # Coluna E = V. CONTÁBIL (1100.00)
    assert float(ws["E4"].value) == 1100.0
    # Coluna F = IPI + DESP (DEVE SER 0.00 ou None, NUNCA 100.00!)
    val_f4 = ws["F4"].value
    assert val_f4 is None or float(val_f4) == 0.0
    wb.close()
