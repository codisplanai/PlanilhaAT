import io
from decimal import Decimal

import openpyxl
from app.core.config import settings
from app.services.excel.template_filler import TemplateFiller
from app.services.supabase_storage import SupabaseStorageService
from app.services.templates_admin.template_manager import TemplateManager


def _tributaria_xlsx_bytes() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Planilha AT"
    ws["A2"] = "CABEÇALHO"
    ws["D3"] = "Nº Nota"
    ws["E3"] = "V. Total"
    ws["H3"] = "MVA"
    ws["I3"] = "Fórmula protegida"
    ws["I4"] = "=E4*2"
    stream = io.BytesIO()
    wb.save(stream)
    wb.close()
    return stream.getvalue()


def test_mapeamento_mva_e_persistido_e_usado_sem_alterar_formula(
    db_session,
    tmp_path,
    monkeypatch,
):
    monkeypatch.setattr(settings, "TEMPLATES_DIR", str(tmp_path))
    monkeypatch.setattr(SupabaseStorageService, "is_configured", lambda: False)

    mapping = {
        "start_row": 4,
        "columns": {
            "numero_nota": "D",
            "v_total": "E",
            "mva": "H",
            "a_dst": "J",
            "a_ori": "K",
        },
        "header_cell": "A2",
        "extra_options": {"aliquota_format": "percent_number"},
    }

    template = TemplateManager.upload_new_template_version(
        db=db_session,
        tipo="antecipacao_tributaria",
        file_bytes=_tributaria_xlsx_bytes(),
        filename="tributaria_com_mva.xlsx",
        mapeamento=mapping,
        capacidade_linhas=100,
        promover_ativo=True,
    )
    assert template.mapeamento_campos["columns"]["mva"] == "H"

    output_path = str(tmp_path / "saida.xlsx")
    TemplateFiller.fill_template(
        template_path=template.arquivo_path,
        mapping=template.mapeamento_campos,
        rows_data=[
            {
                "numero_nota": "123",
                "v_total": Decimal("1000.00"),
                "mva": Decimal("56.75"),
                "a_dst": Decimal("20.50"),
                "a_ori": Decimal("7.00"),
            }
        ],
        output_path=output_path,
        header_info={"razao_social": "PASSO A PASSO CALÇADOS", "competencia": "08/2026"},
    )

    wb = openpyxl.load_workbook(output_path, data_only=False)
    ws = wb["Planilha AT"]
    assert ws["H4"].value == 56.75
    assert ws["I4"].value == "=E4*2"
    wb.close()
