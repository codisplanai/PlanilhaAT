import json
import io
import openpyxl
from datetime import datetime, date
from decimal import Decimal

from app.core.seeds import seed_default_templates

SAMPLE_SPED_E2E = """|0000|019|0|01012026|31012026|EMPRESA TESTE SPED E2E LTDA|12345678000195||BA|123456789|2927408|||A|1|
|0150|FORN01|FORNECEDOR SP LTDA|1058|98765432000180||SP|3550308||RUA TESTE|100||CENTRO|
|0200|PROD01|PRODUTO ELETRONICO A|||UN|01|84713012|||18,00||
|0200|PROD02|PRODUTO ELETRONICO B|||UN|01|84716053|||18,00||
|C100|0|1|FORN01|55|00|1|501|35260198765432000180550010000005011000123456|10012026|20012026|4000,00|0|0,00|0,00|4000,00|0|0,00|0,00|0,00|4000,00|480,00|0,00|0,00|0,00|0,00|0,00|0,00|0,00|
|C170|1|PROD01|PRODUTO ELETRONICO A|1,000|UN|4000,00|0,00|0|000|6102||4000,00|12,00|480,00|0,00|0,00|0,00|0|50|999|0,00|0,00|0,00|
|C100|0|1|FORN01|55|00|1|502|35260198765432000180550010000005021000123456|05012026|12012026|1000,00|0|0,00|0,00|1000,00|0|0,00|0,00|0,00|1000,00|120,00|0,00|0,00|0,00|0,00|0,00|0,00|0,00|
|C170|1|PROD02|PRODUTO ELETRONICO B|1,000|UN|1000,00|0,00|0|000|6102||1000,00|12,00|120,00|0,00|0,00|0,00|0|50|999|0,00|0,00|0,00|
|9999|10|
"""

def test_fluxo_completo_api_via_sped_fiscal(client, create_sample_excel_template):
    # 1. Configurar Perfil de Regras e Empresa
    res_perfil = client.post("/api/v1/perfis-regras", json={
        "nome": "Perfil Teste SPED E2E",
        "configuracoes_extras": {}
    })
    assert res_perfil.status_code in [200, 201]
    perfil_id = res_perfil.json()["id"]

    # Alíquota interna de destino: 20.5% (0.205)
    res_regra = client.post("/api/v1/regras-aliquotas", json={
        "perfil_regras_id": perfil_id,
        "uf": "BA",
        "ncm": None,
        "aliquota": 0.205,
        "descricao": "Alíquota Padrão BA"
    })
    assert res_regra.status_code in [200, 201]

    # Criar Empresa
    res_empresa = client.post("/api/v1/empresas", json={
        "razao_social": "EMPRESA TESTE SPED E2E LTDA",
        "cnpj": "12.345.678/0001-95",
        "uf": "BA",
        "perfil_regras_id": perfil_id,
        "ativo": True
    })
    assert res_empresa.status_code in [200, 201]
    empresa_id = res_empresa.json()["id"]

    # 2. Template XLSX
    template_path = create_sample_excel_template("antecipacao_parcial")
    with open(template_path, "rb") as f:
        template_bytes = f.read()

    mapping = {
        "start_row": 4,
        "sheet_name": "Planilha AT",
        "columns": {
            "numero_nota": "A",
            "data_emissao": "B",
            "ncm": "C",
            "v_total": "D",
            "base_calculo": "E",
            "a_ori": "F",
            "a_dst": "G"
        }
    }

    res_tpl = client.post(
        "/api/v1/templates/upload",
        data={
            "tipo": "antecipacao_parcial",
            "mapeamento_json": json.dumps(mapping),
            "observacoes": "Template Teste SPED",
            "promover_ativo": True
        },
        files={
            "file": ("template_sped.xlsx", template_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        }
    )
    assert res_tpl.status_code in [200, 201]

    # 3. Criar Solicitação
    res_solic = client.post("/api/v1/solicitacoes", json={
        "empresa_id": empresa_id,
        "periodo_inicio": "2026-01-01",
        "periodo_fim": "2026-01-31",
        "tipo_planilha": "antecipacao_parcial"
    })
    assert res_solic.status_code in [200, 201]
    solic_id = res_solic.json()["id"]

    # 4. Processar usando arquivo SPED Fiscal (.txt)
    res_proc = client.post(
        f"/api/v1/solicitacoes/{solic_id}/processar",
        files={
            "sped_file": ("sped_fiscal_202601.txt", SAMPLE_SPED_E2E.encode("utf-8"), "text/plain")
        }
    )
    assert res_proc.status_code == 200, res_proc.text
    dados = res_proc.json()

    assert dados["status"] == "concluido"
    assert dados["total_notas_processadas"] == 2

    # Verificar que foram ordenadas por data de entrada (12/01 para NF 502, 20/01 para NF 501)
    notas_api = dados["notas_processadas"]
    assert [n["numero_nota"] for n in notas_api] == ["502", "501"]
    assert notas_api[0]["data_entrada"] == "2026-01-12"
    assert notas_api[0]["origem_data_entrada"] == "sped_fiscal"
    assert notas_api[1]["data_entrada"] == "2026-01-20"
    assert notas_api[1]["origem_data_entrada"] == "sped_fiscal"

    # Verificar valores calculados para NF 501:
    # V.Total = 4000, A.DST = 20.5% -> Débito = 820.00
    # Base = 4000, A.ORI = 12.0% -> Crédito = 480.00
    # Devido = 820.00 - 480.00 = 340.00
    nf_501 = next(n for n in notas_api if n["numero_nota"] == "501")
    assert float(nf_501["debito"]) == 820.00
    assert float(nf_501["credito"]) == 480.00
    assert float(nf_501["valor_devido"]) == 340.00

    # 5. Baixar e validar a planilha gerada
    res_dl = client.get(f"/api/v1/solicitacoes/{solic_id}/download")
    assert res_dl.status_code == 200

    wb = openpyxl.load_workbook(io.BytesIO(res_dl.content), data_only=False)
    ws = wb["Planilha AT"]

    # Linha 4 -> NF 502
    assert str(ws["A4"].value) == "502"
    assert float(ws["D4"].value) == 1000.00

    # Linha 5 -> NF 501
    assert str(ws["A5"].value) == "501"
    assert float(ws["D5"].value) == 4000.00

    # Fórmulas preservadas nas colunas de cálculo
    assert ws["H4"].value == "=D4*G4"
    assert ws["I4"].value == "=E4*F4"
    assert ws["J4"].value == "=H4-I4"

    wb.close()


def test_sped_fiscal_com_nota_zerada_ou_complementar(client, create_sample_excel_template):
    # 1. Perfil e Empresa
    res_perfil = client.post("/api/v1/perfis-regras", json={
        "nome": "Perfil Teste SPED Zerada",
        "configuracoes_extras": {}
    })
    perfil_id = res_perfil.json()["id"]

    client.post("/api/v1/regras-aliquotas", json={
        "perfil_regras_id": perfil_id,
        "uf": "BA",
        "ncm": None,
        "aliquota": 0.205,
        "descricao": "Padrão"
    })

    res_empresa = client.post("/api/v1/empresas", json={
        "razao_social": "EMPRESA TESTE SPED ZERADA LTDA",
        "cnpj": "12.345.678/0001-95",
        "uf": "BA",
        "perfil_regras_id": perfil_id,
        "ativo": True
    })
    empresa_id = res_empresa.json()["id"]

    # 2. Template
    template_path = create_sample_excel_template("antecipacao_parcial")
    with open(template_path, "rb") as f:
        template_bytes = f.read()

    mapping = {
        "start_row": 4,
        "sheet_name": "Planilha AT",
        "columns": {
            "numero_nota": "A",
            "data_emissao": "B",
            "ncm": "C",
            "v_total": "D",
            "base_calculo": "E",
            "a_ori": "F",
            "a_dst": "G"
        }
    }

    client.post(
        "/api/v1/templates/upload",
        data={
            "tipo": "antecipacao_parcial",
            "mapeamento_json": json.dumps(mapping),
            "observacoes": "Template Teste Zerada",
            "promover_ativo": True
        },
        files={
            "file": ("template_zerada.xlsx", template_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        }
    )

    # 3. SPED contendo 1 nota válida (NF 100) e 1 nota complementar zerada (NF 4600 com v_total = 0.00)
    sped_com_zerada = """|0000|019|0|01012026|31012026|EMPRESA TESTE SPED ZERADA LTDA|12345678000195||BA|123456789|2927408|||A|1|
|0150|FORN01|FORNECEDOR SP LTDA|1058|98765432000180||SP|3550308||RUA TESTE|100||CENTRO|
|0200|PROD01|PRODUTO A|||UN|01|84713012|||18,00||
|C100|0|1|FORN01|55|00|1|100|35260198765432000180550010000001001000123456|10012026|15012026|2000,00|0|0,00|0,00|2000,00|0|0,00|0,00|0,00|2000,00|240,00|0,00|0,00|0,00|0,00|0,00|0,00|0,00|
|C170|1|PROD01|PRODUTO A|1,000|UN|2000,00|0,00|0|000|6102||2000,00|12,00|240,00|0,00|0,00|0,00|0|50|999|0,00|0,00|0,00|
|C100|0|1|FORN01|55|00|100|4600|35260612326987000118551000000046001005354433|19012026|19012026|0,00|0|0,00|0,00|0,00|9|0,00|0,00|0,00|2542,94|521,30|0,00|0,00|0,00|0,00|0,00|0,00|0,00|
|C190|090|6102|20,50|0,00|2542,94|521,30|0,00|0,00|0,00|0,00||
|9999|10|
"""

    res_solic = client.post("/api/v1/solicitacoes", json={
        "empresa_id": empresa_id,
        "periodo_inicio": "2026-01-01",
        "periodo_fim": "2026-01-31",
        "tipo_planilha": "antecipacao_parcial"
    })
    solic_id = res_solic.json()["id"]

    res_proc = client.post(
        f"/api/v1/solicitacoes/{solic_id}/processar",
        files={
            "sped_file": ("sped_fiscal.txt", sped_com_zerada.encode("utf-8"), "text/plain")
        }
    )
    assert res_proc.status_code == 200, res_proc.text
    dados = res_proc.json()

    # O processamento deve ser concluído com sucesso:
    # 1 nota processada (NF 100) e 1 nota ignorada (NF 4600)
    assert dados["status"] == "concluido"
    assert dados["total_notas_processadas"] == 1
    assert len(dados["notas_ignoradas"]) == 1
    assert dados["notas_ignoradas"][0]["numero_nota"] == "4600"
    assert "zerado" in dados["notas_ignoradas"][0]["motivo"].lower()

