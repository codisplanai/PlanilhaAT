import pytest
import os
import io
import openpyxl
from decimal import Decimal
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool
from fastapi.testclient import TestClient

from app.core.database import Base, get_db
from app.core.seeds import seed_default_cfop_rules
from app.main import app
from app.models.perfil_regras import PerfilRegras
from app.models.empresa import Empresa
from app.models.regra_aliquota import RegraAliquotaDestino

# In-memory SQLite compartilhado usando StaticPool para manter tabelas durante todo o teste
SQLALCHEMY_DATABASE_URL = "sqlite://"

engine = create_engine(
    SQLALCHEMY_DATABASE_URL,
    connect_args={"check_same_thread": False},
    poolclass=StaticPool
)
TestingSessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)

@pytest.fixture(scope="function")
def db_session():
    Base.metadata.create_all(bind=engine)
    db = TestingSessionLocal()
    try:
        # Regras padrão globais de roteamento CFOP -> planilha, disponíveis em todo teste
        # (na aplicação real elas são semeadas no startup; o TestClient roda o startup
        # contra o banco de produção, não contra este engine de testes em memória)
        seed_default_cfop_rules(db)
        yield db
    finally:
        db.close()
        Base.metadata.drop_all(bind=engine)

@pytest.fixture(scope="function")
def client(db_session):
    def override_get_db():
        try:
            yield db_session
        finally:
            pass

    app.dependency_overrides[get_db] = override_get_db
    with TestClient(app) as c:
        yield c
    app.dependency_overrides.clear()

@pytest.fixture
def sample_xml_nfe():
    """Gera um XML sintético de NF-e v4.00 com estrutura SEFAZ oficial válida"""
    return b"""<?xml version="1.0" encoding="UTF-8"?>
<nfeProc xmlns="http://www.portalfiscal.inf.br/nfe" versao="4.00">
  <NFe>
    <infNFe Id="NFe35260112345678000195550010000012341000012345" versao="4.00">
      <ide>
        <cUF>35</cUF>
        <cNF>00001234</cNF>
        <natOp>VENDA DE MERCADORIAS</natOp>
        <mod>55</mod>
        <serie>1</serie>
        <nNF>1234</nNF>
        <dhEmi>2026-01-15T10:00:00-03:00</dhEmi>
        <tpNF>1</tpNF>
        <idDest>2</idDest>
      </ide>
      <emit>
        <CNPJ>98765432000180</CNPJ>
        <xNome>FORNECEDOR SAO PAULO LTDA</xNome>
        <enderEmit>
          <UF>SP</UF>
        </enderEmit>
      </emit>
      <dest>
        <CNPJ>12345678000195</CNPJ>
        <xNome>CLIENTE CONTABILIDADE BAHIA LTDA</xNome>
        <enderDest>
          <UF>BA</UF>
        </enderDest>
      </dest>
      <det nItem="1">
        <prod>
          <cProd>PROD001</cProd>
          <xProd>NOTEBOOK CORPORATIVO</xProd>
          <NCM>84713012</NCM>
          <CFOP>6102</CFOP>
          <uCom>UN</uCom>
          <qCom>2.0000</qCom>
          <vUnCom>2500.00</vUnCom>
          <vProd>5000.00</vProd>
          <vFrete>100.00</vFrete>
          <vSeg>0.00</vSeg>
          <vDesc>0.00</vDesc>
          <vOutro>50.00</vOutro>
        </prod>
        <imposto>
          <ICMS>
            <ICMS00>
              <orig>0</orig>
              <CST>00</CST>
              <modBC>3</modBC>
              <vBC>5150.00</vBC>
              <pICMS>12.00</pICMS>
              <vICMS>618.00</vICMS>
            </ICMS00>
          </ICMS>
          <IPI>
            <IPITrib>
              <CST>50</CST>
              <vBC>5000.00</vBC>
              <pIPI>5.00</pIPI>
              <vIPI>250.00</vIPI>
            </IPITrib>
          </IPI>
        </imposto>
      </det>
      <total>
        <ICMSTot>
          <vBC>5150.00</vBC>
          <vICMS>618.00</vICMS>
          <vProd>5000.00</vProd>
          <vFrete>100.00</vFrete>
          <vSeg>0.00</vSeg>
          <vDesc>0.00</vDesc>
          <vOutro>50.00</vOutro>
          <vIPI>250.00</vIPI>
          <vNF>5400.00</vNF>
        </ICMSTot>
      </total>
    </infNFe>
  </NFe>
</nfeProc>"""

@pytest.fixture
def create_sample_excel_template(tmp_path):
    """Cria uma planilha Excel com fórmulas para testar preenchimento e preservação de fórmulas"""
    def _create(tipo="antecipacao_parcial"):
        file_path = str(tmp_path / f"template_{tipo}_test.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Planilha AT"

        # Cabeçalhos
        ws["A1"] = "RELATÓRIO DE ANTECIPAÇÃO PARCIAL"
        ws["A3"] = "Nº Nota"
        ws["B3"] = "Data Emissão"
        ws["C3"] = "NCM"
        ws["D3"] = "V. Total"
        ws["E3"] = "Base Cálculo"
        ws["F3"] = "A. ORI"
        ws["G3"] = "A. DST"
        ws["H3"] = "Débito (Calculado)"
        ws["I3"] = "Crédito (Calculado)"
        ws["J3"] = "Valor Devido (Calculado)"

        # Fórmulas nas linhas 4 a 10 para colunas H, I, J
        for r in range(4, 11):
            ws[f"H{r}"] = f"=D{r}*G{r}"
            ws[f"I{r}"] = f"=E{r}*F{r}"
            ws[f"J{r}"] = f"=H{r}-I{r}"

        # Totais na linha 12 com fórmulas de soma
        ws["C12"] = "TOTALIZADOR:"
        ws["D12"] = "=SUM(D4:D10)"
        ws["E12"] = "=SUM(E4:E10)"
        ws["H12"] = "=SUM(H4:H10)"
        ws["I12"] = "=SUM(I4:I10)"
        ws["J12"] = "=SUM(J4:J10)"

        wb.save(file_path)
        wb.close()
        return file_path

    return _create


def build_xml_nfe(numero: str, chave: str, valor: str, dia: str, cfop: str = "6102", v_ipi: str = "0.00", v_bc: str = None) -> bytes:
    """XML de NF-e mínimo e válido, emitido por fornecedor de SP para cliente da BA em janeiro/2026."""
    bc = v_bc or valor
    ipi_tag = f"<IPI><IPITrib><vIPI>{v_ipi}</vIPI></IPITrib></IPI>" if v_ipi and v_ipi != "0.00" else ""
    return f"""<?xml version="1.0" encoding="UTF-8"?>
<nfeProc xmlns="http://www.portalfiscal.inf.br/nfe">
  <NFe>
    <infNFe Id="NFe{chave}">
      <ide><nNF>{numero}</nNF><serie>1</serie><dhEmi>2026-01-{dia}T10:00:00-03:00</dhEmi></ide>
      <emit><CNPJ>98765432000180</CNPJ><enderEmit><UF>SP</UF></enderEmit></emit>
      <dest><CNPJ>12345678000195</CNPJ><enderDest><UF>BA</UF></enderDest></dest>
      <total><ICMSTot><vNF>{valor}</vNF><vBC>{bc}</vBC></ICMSTot></total>
      <det nItem="1">
        <prod><NCM>21069090</NCM><CFOP>{cfop}</CFOP><vProd>{bc}</vProd>
          <vFrete>0.00</vFrete><vSeg>0.00</vSeg><vOutro>0.00</vOutro><vDesc>0.00</vDesc></prod>
        <imposto><ICMS><ICMS00><vBC>{bc}</vBC><pICMS>12.00</pICMS></ICMS00></ICMS>{ipi_tag}</imposto>
      </det>
    </infNFe>
  </NFe>
</nfeProc>""".encode("utf-8")


CHAVE_NF901 = "35260198765432000180550010000009011000000901"
CHAVE_NF902 = "35260198765432000180550010000009021000000902"
CHAVE_NF903 = "35260198765432000180550010000009031000000903"

# SPED de janeiro/2026 contendo APENAS a NF 901: é a única cuja mercadoria entrou no mês.
_SPED_JANEIRO_COM_NF901 = """|0000|019|0|01012026|31012026|Cliente BA|12345678000195||BA|123|2927408|||A|1|
|0150|F1|FORNECEDOR SP|1058|98765432000180||SP|3550308||R|1||C|
|0200|P1|PRODUTO|||UN|01|21069090|||18,00||
|C100|0|1|F1|55|00|1|901|35260198765432000180550010000009011000000901|10012026|20012026|2000,00|0|0,00|0,00|2000,00|0|0,00|0,00|0,00|2000,00|240,00|0,00|0,00|0,00|0,00|0,00|0,00|0,00|
|C170|1|P1|PRODUTO|1,000|UN|2000,00|0,00|0|000|6102||2000,00|12,00|240,00|0,00|0,00|0,00|0|50|999|0,00|0,00|0,00|
|9999|5|
"""


@pytest.fixture
def sped_janeiro_com_nf901() -> bytes:
    return _SPED_JANEIRO_COM_NF901.encode("utf-8")


@pytest.fixture
def xml_nf901() -> bytes:
    return build_xml_nfe("901", CHAVE_NF901, "2000.00", "10")


@pytest.fixture
def xml_nf902() -> bytes:
    return build_xml_nfe("902", CHAVE_NF902, "3000.00", "28")


@pytest.fixture
def xml_nf903_difal() -> bytes:
    return build_xml_nfe("903", CHAVE_NF903, "500.00", "29", cfop="6556")


@pytest.fixture
def cenario_janeiro(db_session, create_sample_excel_template):
    """
    Monta perfil de regras, empresa da BA, alíquota padrão 18% e uma solicitação de
    janeiro/2026 em modo automático (tipo_planilha=None), com templates ativos para
    os tipos pedidos. Devolve a Solicitacao pronta para processar.
    """
    from decimal import Decimal
    from datetime import date
    from app.models.perfil_regras import PerfilRegras
    from app.models.empresa import Empresa
    from app.models.regra_aliquota import RegraAliquotaDestino
    from app.models.solicitacao import Solicitacao
    from app.services.templates_admin.template_manager import TemplateManager

    def _montar(tipos=("antecipacao_parcial", "antecipacao_parcial_antecipado")):
        for tipo in tipos:
            path = create_sample_excel_template(tipo=tipo)
            with open(path, "rb") as f:
                TemplateManager.upload_new_template_version(
                    db=db_session, tipo=tipo, filename=f"t_{tipo}.xlsx", file_bytes=f.read(),
                    mapeamento={"start_row": 4, "columns": {"numero_nota": "A", "v_total": "D"}},
                    promover_ativo=True,
                )

        perfil = PerfilRegras(nome="Perfil Cenario Janeiro")
        db_session.add(perfil)
        db_session.commit()
        db_session.add(RegraAliquotaDestino(
            perfil_regras_id=perfil.id, uf="BA", ncm=None, aliquota=Decimal("0.1800")))
        db_session.commit()

        empresa = Empresa(razao_social="Cliente BA", cnpj="12345678000195", uf="BA",
                          perfil_regras_id=perfil.id)
        db_session.add(empresa)
        db_session.commit()

        sol = Solicitacao(empresa_id=empresa.id, periodo_inicio=date(2026, 1, 1),
                          periodo_fim=date(2026, 1, 31), tipo_planilha=None, status="pendente")
        db_session.add(sol)
        db_session.commit()
        return sol

    return _montar
