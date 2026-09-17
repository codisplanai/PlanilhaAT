import io
import os
import re
import datetime
from typing import Optional, List, Dict, Tuple, Set
from dataclasses import dataclass
import openpyxl
import xlrd

from app.core.exceptions import ValidationException

@dataclass
class PlanilhaEntradaRecord:
    numero_raw: str
    numero_normalizado: str
    serie_normalizada: str
    cnpj_emitente_normalizado: str
    chave_acesso_normalizada: str
    data_entrada: Optional[datetime.date]
    cfop_normalizado: str = ""

class DataEntradaNormalizer:
    """
    Funções puras de normalização obrigatória para cruzamento de dados contábeis.
    """

    @staticmethod
    def normalize_cnpj(cnpj: Optional[str]) -> str:
        """Mantém apenas os dígitos do CNPJ/CPF"""
        if not cnpj:
            return ""
        return re.sub(r"\D", "", DataEntradaNormalizer._numeric_text(cnpj))

    @staticmethod
    def normalize_numero(numero: Optional[str]) -> str:
        """Remove zeros à esquerda e caracteres não numéricos do número da nota"""
        if not numero:
            return ""
        clean = re.sub(r"\D", "", DataEntradaNormalizer._numeric_text(numero))
        return clean.lstrip("0")

    @staticmethod
    def normalize_serie(serie: Optional[str]) -> str:
        """Remove espaços em branco da série da nota"""
        if serie is None:
            return ""
        clean = DataEntradaNormalizer._numeric_text(serie)
        # Se vier com zeros à esquerda ex: '001', manter dígitos limpos
        return clean

    @staticmethod
    def normalize_chave(chave: Optional[str]) -> str:
        """Mantém apenas dígitos da chave de acesso"""
        if not chave:
            return ""
        clean = re.sub(r"\D", "", DataEntradaNormalizer._numeric_text(chave))
        return clean if len(clean) == 44 else ""

    @staticmethod
    def normalize_cfop(cfop: Optional[str]) -> str:
        """Mantém apenas dígitos do CFOP"""
        if not cfop:
            return ""
        return re.sub(r"\D", "", DataEntradaNormalizer._numeric_text(cfop))

    @staticmethod
    def _numeric_text(value: object) -> str:
        text = str(value).strip()
        return text[:-2] if re.fullmatch(r"\d+\.0", text) else text


class PlanilhaEntradaParser:
    """
    Parser da planilha de exportação do sistema contábil (suporta layout Prosoft e similares em .xls/.xlsx).
    """

    @staticmethod
    def _find_column_index(headers: List[str], exact_names: List[str], partial_names: Optional[List[str]] = None) -> Optional[int]:
        # 1. Busca por nome exato (case-insensitive)
        for exact in exact_names:
            exact_clean = exact.strip().lower()
            for idx, h in enumerate(headers):
                if h.strip().lower() == exact_clean:
                    return idx
        # 2. Busca parcial (excluindo falsos positivos como 'IE do Terceiro')
        if partial_names:
            for partial in partial_names:
                partial_clean = partial.strip().lower()
                for idx, h in enumerate(headers):
                    h_clean = h.strip().lower()
                    if partial_clean in h_clean and "ie " not in h_clean and "inscrição" not in h_clean:
                        return idx
        return None

    @classmethod
    def parse(cls, file_bytes: bytes, filename: str = "planilha.xlsx") -> List[PlanilhaEntradaRecord]:
        if not file_bytes:
            return []

        suffix = os.path.splitext(filename)[1].lower()
        try:
            if suffix == ".xlsx":
                return cls._parse_openpyxl(file_bytes)
            if suffix == ".xls":
                return cls._parse_xlrd(file_bytes)
            raise ValidationException("A planilha contábil deve possuir extensão .xls ou .xlsx.")
        except ValidationException:
            raise
        except Exception as exc:
            raise ValidationException(
                f"Não foi possível ler a planilha do sistema contábil '{filename}'. "
                "Verifique se o arquivo é um Excel válido."
            ) from exc

    @classmethod
    def _parse_openpyxl(cls, file_bytes: bytes) -> List[PlanilhaEntradaRecord]:
        bio = io.BytesIO(file_bytes)
        wb = openpyxl.load_workbook(bio, data_only=True)
        ws = wb.active

        # Localizar linha de cabeçalho (busca nas primeiras 15 linhas)
        header_row_idx = None
        header_names: List[str] = []

        for r in range(1, min(16, ws.max_row + 1)):
            row_vals = [str(ws.cell(r, c).value or "").strip() for c in range(1, ws.max_column + 1)]
            # Critério: linha que contenha 'Número Nota' ou 'Dt.Escritur'
            if any("número nota" in v.lower() or "numero nota" in v.lower() or "dt.escritur" in v.lower() for v in row_vals):
                header_row_idx = r
                header_names = row_vals
                break

        if header_row_idx is None:
            # Se não encontrou cabeçalho específico, usa a linha 1 como cabeçalho
            header_row_idx = 1
            header_names = [str(ws.cell(1, c).value or "").strip() for c in range(1, ws.max_column + 1)]

        # Resolver índices das colunas (0-based)
        col_num_idx = cls._find_column_index(header_names, ["Número Nota", "Numero Nota", "Nº Nota", "Num Nota", "N. Fiscal", "Nota"])
        col_dt_idx = cls._find_column_index(header_names, ["Dt.Escritur.", "Dt.Escritur", "Data Escrituração", "Data Entrada", "Dt.Entrada", "Data de Entrada", "Data Entrada/Saída"])
        col_serie_idx = cls._find_column_index(header_names, ["Série", "Serie", "Ser"])
        col_cnpj_idx = cls._find_column_index(header_names, ["Terceiro", "CNPJ do Emitente", "CNPJ Emitente", "CNPJ/CPF", "CNPJ", "CPF/CNPJ"])
        col_chave_idx = cls._find_column_index(header_names, ["Chave da Nota Fiscal Eletrônica", "Chave NFe", "Chave de Acesso", "Chave Eletrônica", "Chave"])
        col_cfop_idx = cls._find_column_index(
            header_names,
            ["CFOP", "C.F.O.P.", "Cód. Fiscal", "Cod. Fiscal", "Código Fiscal", "Natureza da Operação", "Natureza"],
            partial_names=["cfop"]
        )

        if col_num_idx is None or col_dt_idx is None:
            wb.close()
            raise ValidationException(
                "A planilha contábil precisa conter as colunas de número da nota e data de entrada/escrituração."
            )

        records: List[PlanilhaEntradaRecord] = []

        for r in range(header_row_idx + 1, ws.max_row + 1):
            num_val = ws.cell(r, col_num_idx + 1).value if col_num_idx is not None else None
            if num_val is None or str(num_val).strip() == "":
                continue

            num_str = str(num_val).strip()
            dt_raw = ws.cell(r, col_dt_idx + 1).value if col_dt_idx is not None else None
            serie_raw = ws.cell(r, col_serie_idx + 1).value if col_serie_idx is not None else ""
            cnpj_raw = ws.cell(r, col_cnpj_idx + 1).value if col_cnpj_idx is not None else ""
            chave_raw = ws.cell(r, col_chave_idx + 1).value if col_chave_idx is not None else ""
            cfop_raw = ws.cell(r, col_cfop_idx + 1).value if col_cfop_idx is not None else ""

            # Parse de data
            parsed_date: Optional[datetime.date] = None
            if isinstance(dt_raw, datetime.datetime):
                parsed_date = dt_raw.date()
            elif isinstance(dt_raw, datetime.date):
                parsed_date = dt_raw
            elif isinstance(dt_raw, str) and dt_raw.strip():
                clean_dt_str = dt_raw.strip()
                for fmt in ["%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%y"]:
                    try:
                        parsed_date = datetime.datetime.strptime(clean_dt_str, fmt).date()
                        break
                    except ValueError:
                        pass

            rec = PlanilhaEntradaRecord(
                numero_raw=num_str,
                numero_normalizado=DataEntradaNormalizer.normalize_numero(num_str),
                serie_normalizada=DataEntradaNormalizer.normalize_serie(serie_raw),
                cnpj_emitente_normalizado=DataEntradaNormalizer.normalize_cnpj(cnpj_raw),
                chave_acesso_normalizada=DataEntradaNormalizer.normalize_chave(chave_raw),
                data_entrada=parsed_date,
                cfop_normalizado=DataEntradaNormalizer.normalize_cfop(cfop_raw)
            )
            records.append(rec)

        wb.close()
        return records

    @classmethod
    def _parse_xlrd(cls, file_bytes: bytes) -> List[PlanilhaEntradaRecord]:
        book = xlrd.open_workbook(file_contents=file_bytes)
        sheet = book.sheet_by_index(0)

        header_row_idx = None
        header_names: List[str] = []

        for r in range(min(16, sheet.nrows)):
            row_vals = [str(sheet.cell_value(r, c) or "").strip() for c in range(sheet.ncols)]
            if any("número nota" in v.lower() or "numero nota" in v.lower() or "dt.escritur" in v.lower() for v in row_vals):
                header_row_idx = r
                header_names = row_vals
                break

        if header_row_idx is None:
            header_row_idx = 0
            header_names = [str(sheet.cell_value(0, c) or "").strip() for c in range(sheet.ncols)]

        col_num_idx = cls._find_column_index(header_names, ["Número Nota", "Numero Nota", "Nº Nota", "Num Nota", "N. Fiscal", "Nota"])
        col_dt_idx = cls._find_column_index(header_names, ["Dt.Escritur.", "Dt.Escritur", "Data Escrituração", "Data Entrada", "Dt.Entrada", "Data de Entrada"])
        col_serie_idx = cls._find_column_index(header_names, ["Série", "Serie", "Ser"])
        col_cnpj_idx = cls._find_column_index(header_names, ["Terceiro", "CNPJ do Emitente", "CNPJ Emitente", "CNPJ/CPF", "CNPJ"])
        col_chave_idx = cls._find_column_index(header_names, ["Chave da Nota Fiscal Eletrônica", "Chave NFe", "Chave de Acesso", "Chave"])
        col_cfop_idx = cls._find_column_index(
            header_names,
            ["CFOP", "C.F.O.P.", "Cód. Fiscal", "Cod. Fiscal", "Código Fiscal", "Natureza da Operação", "Natureza"],
            partial_names=["cfop"]
        )

        if col_num_idx is None or col_dt_idx is None:
            raise ValidationException(
                "A planilha contábil precisa conter as colunas de número da nota e data de entrada/escrituração."
            )

        records: List[PlanilhaEntradaRecord] = []

        for r in range(header_row_idx + 1, sheet.nrows):
            num_val = sheet.cell_value(r, col_num_idx) if col_num_idx is not None else None
            if num_val is None or str(num_val).strip() == "":
                continue

            num_str = str(num_val).strip()
            # Remover .0 se veio como float no xlrd
            if num_str.endswith(".0"):
                num_str = num_str[:-2]

            dt_cell = sheet.cell(r, col_dt_idx) if col_dt_idx is not None else None
            serie_raw = sheet.cell_value(r, col_serie_idx) if col_serie_idx is not None else ""
            cnpj_raw = sheet.cell_value(r, col_cnpj_idx) if col_cnpj_idx is not None else ""
            chave_raw = sheet.cell_value(r, col_chave_idx) if col_chave_idx is not None else ""
            cfop_raw = sheet.cell_value(r, col_cfop_idx) if col_cfop_idx is not None else ""

            parsed_date: Optional[datetime.date] = None
            if dt_cell is not None:
                if dt_cell.ctype == xlrd.XL_CELL_DATE:
                    date_tuple = xlrd.xldate_as_tuple(dt_cell.value, book.datemode)
                    parsed_date = datetime.date(date_tuple[0], date_tuple[1], date_tuple[2])
                elif isinstance(dt_cell.value, str) and dt_cell.value.strip():
                    for fmt in ["%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"]:
                        try:
                            parsed_date = datetime.datetime.strptime(dt_cell.value.strip(), fmt).date()
                            break
                        except ValueError:
                            pass

            rec = PlanilhaEntradaRecord(
                numero_raw=num_str,
                numero_normalizado=DataEntradaNormalizer.normalize_numero(num_str),
                serie_normalizada=DataEntradaNormalizer.normalize_serie(serie_raw),
                cnpj_emitente_normalizado=DataEntradaNormalizer.normalize_cnpj(cnpj_raw),
                chave_acesso_normalizada=DataEntradaNormalizer.normalize_chave(chave_raw),
                data_entrada=parsed_date,
                cfop_normalizado=DataEntradaNormalizer.normalize_cfop(cfop_raw)
            )
            records.append(rec)

        return records


class DataEntradaMatcher:
    """
    Motor determinístico de correspondência de Data de Entrada e CFOP Contábil:
    - Prioridade 1: Chave de Acesso (44 dígitos) presente em ambos os lados.
    - Prioridade 2: CNPJ Emitente + Série + Número da Nota (todos normalizados).
    
    Regra de Não-Ambiguidade:
    - Se houver exatamente 1 correspondência inequívoca (ou múltiplas linhas com a mesma data exata) -> Preenche com origem 'planilha_sistema_contabil'.
    - Se houver 0 correspondências ou mais de uma data conflitante -> Retorna (None, None) para preenchimento manual posterior.
    """

    @classmethod
    def match_records(
        cls,
        nf_chave: Optional[str],
        nf_cnpj_emitente: Optional[str],
        nf_serie: Optional[str],
        nf_numero: Optional[str],
        planilha_records: List[PlanilhaEntradaRecord]
    ) -> List[PlanilhaEntradaRecord]:
        if not planilha_records:
            return []

        # 1. Tentativa por Chave de Acesso (Prioridade 1)
        norm_chave = DataEntradaNormalizer.normalize_chave(nf_chave)
        if norm_chave and len(norm_chave) == 44:
            matched_by_chave = [r for r in planilha_records if r.chave_acesso_normalizada == norm_chave]
            if matched_by_chave:
                return matched_by_chave

        # 2. Tentativa por CNPJ Emitente + Série + Número (Prioridade 2)
        norm_cnpj = DataEntradaNormalizer.normalize_cnpj(nf_cnpj_emitente)
        norm_serie = DataEntradaNormalizer.normalize_serie(nf_serie)
        norm_num = DataEntradaNormalizer.normalize_numero(nf_numero)

        if norm_cnpj and norm_num:
            serie_variantes = {norm_serie, norm_serie.lstrip("0")} if norm_serie else {"", "0", "001", "1"}
            matched_by_tuple = [
                r for r in planilha_records
                if r.cnpj_emitente_normalizado == norm_cnpj
                and r.numero_normalizado == norm_num
                and (
                    r.serie_normalizada in serie_variantes
                    or (not norm_serie)
                    or (not r.serie_normalizada)
                    or r.serie_normalizada.lstrip("0") == norm_serie.lstrip("0")
                )
            ]
            if matched_by_tuple:
                return matched_by_tuple

        return []

    @classmethod
    def match_data_entrada(
        cls,
        nf_chave: Optional[str],
        nf_cnpj_emitente: Optional[str],
        nf_serie: Optional[str],
        nf_numero: Optional[str],
        planilha_records: List[PlanilhaEntradaRecord]
    ) -> Tuple[Optional[datetime.date], Optional[str]]:
        matched = cls.match_records(nf_chave, nf_cnpj_emitente, nf_serie, nf_numero, planilha_records)
        if not matched:
            return None, None

        dates_set: Set[datetime.date] = {r.data_entrada for r in matched if r.data_entrada is not None}
        if len(dates_set) == 1:
            return dates_set.pop(), "planilha_sistema_contabil"
        return None, None

    @classmethod
    def match_cfop(
        cls,
        nf_chave: Optional[str],
        nf_cnpj_emitente: Optional[str],
        nf_serie: Optional[str],
        nf_numero: Optional[str],
        planilha_records: List[PlanilhaEntradaRecord]
    ) -> Optional[str]:
        """Retorna o CFOP atribuído à nota na planilha contábil se houver valor único e não ambíguo."""
        matched = cls.match_records(nf_chave, nf_cnpj_emitente, nf_serie, nf_numero, planilha_records)
        if not matched:
            return None

        cfops_set = {r.cfop_normalizado for r in matched if r.cfop_normalizado}
        if len(cfops_set) == 1:
            return cfops_set.pop()
        return None
