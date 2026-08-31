from typing import Dict, Tuple, Set
import openpyxl
from openpyxl.worksheet.worksheet import Worksheet
from app.core.exceptions import TemplateIntegrityException

class FormulaGuard:
    """
    Guardião de fórmulas de planilhas Excel.
    Garante que células com fórmulas nunca sejam sobrescritas durante o preenchimento de dados
    e verifica a integridade absoluta das fórmulas antes e depois da escrita.
    """

    @classmethod
    def extract_formula_map(cls, wb: openpyxl.Workbook) -> Dict[str, Dict[Tuple[int, int], str]]:
        """
        Mapeia todas as fórmulas existentes em cada aba da planilha.
        Retorna: { sheet_name: { (row, col): formula_string } }
        """
        formula_map: Dict[str, Dict[Tuple[int, int], str]] = {}

        for sheet in wb.worksheets:
            sheet_formulas: Dict[Tuple[int, int], str] = {}
            for row in sheet.iter_rows():
                for cell in row:
                    val = cell.value
                    if val is not None and isinstance(val, str) and val.startswith("="):
                        sheet_formulas[(cell.row, cell.column)] = val
                    elif cell.data_type == "f":
                        sheet_formulas[(cell.row, cell.column)] = str(val)
            formula_map[sheet.title] = sheet_formulas

        return formula_map

    @classmethod
    def assert_no_formula_overwrite(
        cls,
        sheet: Worksheet,
        row_idx: int,
        col_idx: int,
        field_name: str
    ) -> None:
        """
        Verifica se a célula onde se pretende escrever contém uma fórmula.
        Se contiver, bloqueia a escrita e lança TemplateIntegrityException.
        """
        cell = sheet.cell(row=row_idx, column=col_idx)
        val = cell.value
        if (val is not None and isinstance(val, str) and val.startswith("=")) or cell.data_type == "f":
            col_letter = openpyxl.utils.get_column_letter(col_idx)
            raise TemplateIntegrityException(
                f"Tentativa de sobrescrever fórmula na célula {col_letter}{row_idx} (Aba: '{sheet.title}') "
                f"com o campo '{field_name}'. Conteúdo original da célula: '{val}'. "
                f"O preenchimento foi cancelado para preservar as fórmulas originais."
            )

    @classmethod
    def verify_wb_integrity(
        cls,
        original_formula_map: Dict[str, Dict[Tuple[int, int], str]],
        filled_wb: openpyxl.Workbook
    ) -> None:
        """
        Compara o mapa original de fórmulas com o workbook preenchido para certificar
        que 100% das fórmulas originais da(s) aba(s) preenchida(s) continuam presentes e inalteradas.
        """
        filled_map = cls.extract_formula_map(filled_wb)
        if not filled_map:
            raise TemplateIntegrityException("Nenhuma aba encontrada na planilha preenchida.")

        for sheet_name, filled_formulas in filled_map.items():
            orig_formulas = original_formula_map.get(sheet_name, {})

            for (row, col), orig_formula in orig_formulas.items():
                col_letter = openpyxl.utils.get_column_letter(col)
                if (row, col) not in filled_formulas:
                    raise TemplateIntegrityException(
                        f"Fórmula na célula {col_letter}{row} da aba '{sheet_name}' foi perdida ou apagada! "
                        f"Fórmula esperada: '{orig_formula}'."
                    )

                curr_formula = filled_formulas[(row, col)]
                if curr_formula != orig_formula:
                    raise TemplateIntegrityException(
                        f"Fórmula na célula {col_letter}{row} da aba '{sheet_name}' foi alterada! "
                        f"Esperada: '{orig_formula}', Atual: '{curr_formula}'."
                    )
