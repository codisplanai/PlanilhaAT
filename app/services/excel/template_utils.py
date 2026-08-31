"""Seleção de abas e conversão de valores para preenchimento de templates."""

from datetime import date, datetime
from decimal import Decimal
from typing import Any, Mapping, Optional, Sequence

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet


MONTH_NAMES = (
    "JANEIRO",
    "FEVEREIRO",
    "MARÇO",
    "ABRIL",
    "MAIO",
    "JUNHO",
    "JULHO",
    "AGOSTO",
    "SETEMBRO",
    "OUTUBRO",
    "NOVEMBRO",
    "DEZEMBRO",
)

PERCENTAGE_FIELDS = frozenset({"a_dst", "a_ori", "mva", "reducao"})


def select_worksheet(
    workbook: Workbook,
    mapping: Mapping[str, Any],
    rows: Sequence[Mapping[str, Any]],
    header_info: Optional[Mapping[str, Any]],
) -> tuple[Worksheet, Optional[int], Optional[int]]:
    sheet_name = mapping.get("sheet_name")
    month, year = _resolve_period(rows, header_info)

    candidates: list[str] = []
    if sheet_name and str(sheet_name).lower() != "auto":
        candidates.append(str(sheet_name))

    if month and 1 <= month <= 12:
        candidates.append(MONTH_NAMES[month - 1])
        if year:
            candidates.extend(
                [
                    f"{month:02d}-{year}",
                    f"{month:02d}/{year}",
                    f"{month:02d}_{year}",
                    f"{month:02d}-{year} FCP",
                    f"{month:02d}-{year} FCP ",
                ]
            )

    normalized_names = {name.strip().upper(): name for name in workbook.sheetnames}
    for candidate in candidates:
        actual_name = normalized_names.get(candidate.strip().upper())
        if actual_name:
            return workbook[actual_name], month, year

    if month and year:
        prefix = f"{month:02d}-{year}".upper()
        for actual_name in workbook.sheetnames:
            if actual_name.strip().upper().startswith(prefix):
                return workbook[actual_name], month, year

    if sheet_name and sheet_name in workbook.sheetnames:
        return workbook[sheet_name], month, year
    return workbook.active, month, year


def build_header_text(header_info: Mapping[str, Any]) -> str:
    company = str(header_info.get("razao_social") or "").strip()
    period = str(header_info.get("competencia") or "").strip()
    state_registration = str(
        header_info.get("ie") or header_info.get("inscricao_estadual") or ""
    ).strip()

    if state_registration:
        return (
            f"                  Empresa -{company}"
            f"                                       IE: {state_registration}"
            f"                             COMP.  {period}"
        )
    return (
        f"                  Empresa -{company}"
        f"                                                                    COMP.  {period}"
    )


def format_excel_value(field: str, raw_value: Any, percentage_format: str) -> Any:
    if field in PERCENTAGE_FIELDS:
        if percentage_format == "percent_number":
            if isinstance(raw_value, (Decimal, float)):
                value = Decimal(str(raw_value))
                return float(value * Decimal("100.0")) if value <= 1 else float(value)
        elif isinstance(raw_value, Decimal):
            return float(raw_value)
        return raw_value

    if isinstance(raw_value, Decimal):
        return float(raw_value)
    if isinstance(raw_value, date):
        return datetime(raw_value.year, raw_value.month, raw_value.day)
    return raw_value


def _resolve_period(
    rows: Sequence[Mapping[str, Any]],
    header_info: Optional[Mapping[str, Any]],
) -> tuple[Optional[int], Optional[int]]:
    if header_info:
        raw_month = header_info.get("mes")
        raw_year = header_info.get("ano")
        month: Optional[int] = None
        year: Optional[int] = None

        if raw_month is not None:
            try:
                month = int(raw_month)
            except (ValueError, TypeError):
                month = None

        if raw_year is not None:
            try:
                year = int(raw_year)
            except (ValueError, TypeError):
                year = None

        if "competencia" in header_info and (month is None or year is None):
            parts = str(header_info["competencia"]).strip().split("/")
            if len(parts) == 2:
                if month is None and parts[0].isdigit():
                    month = int(parts[0])
                if year is None and parts[1].isdigit():
                    year = int(parts[1])

        return month, year

    if rows:
        issue_date = rows[0].get("data_emissao")
        if isinstance(issue_date, (datetime, date)):
            return issue_date.month, issue_date.year
    return None, None
